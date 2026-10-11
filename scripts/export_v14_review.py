"""Human-readable v1.4 review workbook and delivery copies; no classifier changes."""
from pathlib import Path
import csv,json,shutil,zipfile
import pyarrow.parquet as pq
from openpyxl import Workbook,load_workbook
from kyiv_alerts.controlled_review.workbook import write_sheet,dropdown
ROOT=Path(__file__).resolve().parents[1];REPORT=ROOT/'reports/classification_v1_4';DATA=ROOT/'data/processed/classification_v1_4'
DEST=ROOT.parent/'downloads';DEST.mkdir(exist_ok=True)
PUBLISH=ROOT/'review_downloads/task0046';PUBLISH.mkdir(parents=True,exist_ok=True)
wb=Workbook();wb.remove(wb.active)
r=json.loads((REPORT/'classification_comparison.json').read_text())
intro=[{'topic':'Статус','description':'EXPERIMENTAL — HUMAN VALIDATION REQUIRED'},
 {'topic':'Метод','description':'v1.3 збережено. v1.4 додає експериментальну атрибуцію за контекстом, короткі епізоди і 10 хвилин перед початком тривоги. Бандероль = реактивний БПЛА за уточненням керівника дослідження. Тип моделі не підтверджується.'},
 {'topic':'Результат','description':f"437 → {r['v1_4']['classified_alerts']} тривог; +{r['comparison']['gained_count']}, втрат {r['comparison']['lost_count']}. Покриття правилами не є точністю."},
 {'topic':'Контроль 2581','description':'За рішенням керівника залишено розбіжність: джерела підтримують jet_uav; контрольна мітка також містить shahed_uav. Не створено правило за ID.'},
 {'topic':'Контроль 2590','description':'shahed_uav — робоча гіпотеза, не незалежно підтверджена людська мітка.'},
 {'topic':'Внески','description':'Rule Gains показує перехресні підтримки та окреме вимкнення кожного правила. Результати не підсумовувати.'},
 {'topic':'Перевірка','description':'Review Sample містить оригінальні тексти, ID, час, буквальну географію та причини невизначеності. Human decision і notes залишено порожніми.'},
 {'topic':'Повні дані','description':'Parquet і звіти збережено окремо від v1.3. Ця книга містить усі тривоги та вибірку джерел; повні маркери є в Parquet. Немає прогнозів чи аналізу часових рядів.'}]
ws=write_sheet(wb,'README',intro);ws.column_dimensions['B'].width=120
f=pq.read_table(DATA/'alert_threat_features.parquet').to_pylist();overview=[]
for x in f:
 overview.append({k:x[k] for k in ['alert_id','start_time','end_time','classification_status','baseline_classification_status','baseline_categories','threat_combination','supporting_message_ids','has_pre_alert_evidence','evidence_strengths','correction_causes','uncertainty_marker_count']}|dict(human_decision=None,human_notes=None))
ws=write_sheet(wb,'All Alerts',overview);dropdown(ws,'human_decision',['confirm_context','reject_context','insufficient_information','needs_discussion'])
for title,filename in [('Reviewed Cases','reviewed_case_checks.csv'),('New Alerts','newly_classified_alerts.csv'),('Rule Gains','rule_gain_breakdown.csv'),('Rule Ablations','rule_ablation_comparison.csv'),('Months','monthly_coverage.csv'),('Years','yearly_coverage.csv'),('Review Sample','review_sample.csv')]:
 with (REPORT/filename).open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.DictReader(stream))
 if title=='Review Sample':
  for row in rows:
   row['human_decision']=None;row['human_notes']=None
 ws=write_sheet(wb,title,rows)
 if title=='Review Sample':dropdown(ws,'human_decision',['confirm_context','reject_context','insufficient_information','needs_discussion'])
path=REPORT/'classification_review.xlsx';wb.save(path);loaded=load_workbook(path);assert loaded['All Alerts'].max_row==len(f)+1;loaded.close()
shutil.copy2(path,DEST/'task0046_classification_review.xlsx');shutil.copy2(path,PUBLISH/'classification_review.xlsx')
# Git delivery folder contains review artifacts only; protected datasets stay isolated.
for p in sorted(REPORT.iterdir()):
 if p.suffix in {'.csv','.md','.json'}:shutil.copy2(p,PUBLISH/p.name)
with zipfile.ZipFile(DEST/'task0046_classification_bundle.zip','w',zipfile.ZIP_DEFLATED) as z:
 for folder in [REPORT,DATA]:
  for p in sorted(folder.iterdir()):z.write(p,p.relative_to(ROOT))
 for name in ['attribution_v1_4.yaml','task0046_reviewed_cases.yaml']:z.write(ROOT/'config'/name,'config/'+name)
print(path)
