"""Diagnostic signals must not become labels; full evidence cannot be omitted."""
import copy
import csv
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pyarrow.parquet as pq
import pytest
import yaml
from openpyxl import load_workbook

from kyiv_alerts.residual_review.core import (gap_hits, segment, select_residual,
    QUOTAS, HUMAN_FIELDS, message_signals)
from kyiv_alerts.residual_review.__main__ import (FILES, OUTPUT, baseline, snapshot,
    window, diagnostic_readouts, validate_export)
from kyiv_alerts.classification_v14.lexical import LexicalProfile

ROOT = Path(__file__).resolve().parents[1]
T = datetime(2026, 7, 1, 10, tzinfo=timezone.utc)
PROFILE = LexicalProfile(ROOT/'config/threat_dictionary_v1_3.yaml', ROOT/'config/kyiv_geography_v1_1.yaml',
                        yaml.safe_load((ROOT/'config/attribution_v1_4.yaml').read_text()))


def msg(text, mid='m', delta=0):
    return dict(message_id=mid, text=text, message_time=T+timedelta(minutes=delta))


def feature():
    return dict(alert_id='synthetic', start_time=T, classification_status='no_marker_with_text')


def test_diagnostic_gap_preserves_span_without_assigning_category():
    text='На Київ реактив берковець'
    hits=gap_hits(text, [])
    assert hits and all(text[h['match_start']:h['match_end']]==h['matched_text'] for h in hits)
    assert all('category' not in h and h['interpretation']=='diagnostic_only_no_category' for h in hits)
    h=hits[0]
    assert not gap_hits(text, [dict(evidence_role='lexical_anchor', match_start=h['match_start'], match_end=h['match_end'])])


@pytest.mark.parametrize('text', ['Жарт: реактивний мопед', 'Збір на мопед', 'Не фіксується реактив', 'Бандероль?'])
def test_nonoperational_probes_do_not_create_gap_signal(text):
    assert not gap_hits(text, [])


def test_place_only_sparse_text_is_b_and_f_without_category():
    r=segment(feature(), [msg('Оболонь')], [], PROFILE, [])
    assert r['diagnostic_segments']==['B','F'] and r['primary_segment']=='B'
    assert 'category' not in r and r['rejected_in_alert_lexical_count']==0


def test_rejected_lexical_takes_primary_priority_over_gap():
    marker=dict(message_id='m', evidence_role='lexical_anchor', time_relation='inside_interval',
                match_start=0, match_end=4, accepted_for_aggregation=False, exclusion_reason='no_explicit_current_operational_signal', rule_id='frozen.test')
    r=segment(feature(), [msg('БпЛА реактив')], [marker], PROFILE, [])
    assert set(r['diagnostic_segments'])=={'A','C'} and r['primary_segment']=='A'
    assert r['recorded_rejection_reasons']=={'no_explicit_current_operational_signal':1}


def test_nonop_majority_is_proxy_and_can_overlap_insufficient():
    r=segment(feature(), [msg('Вчора було'),msg('Жарт',mid='n')], [], PROFILE, [])
    assert r['primary_segment']=='E' and r['diagnostic_segments']==['E','F']
    assert r['source_information_status']=='diagnostic_not_validated'


def test_post_boundary_hit_is_diagnostic_not_accepted():
    read=diagnostic_readouts([msg('Київ БпЛА над центром')], PROFILE, {})
    assert read and all(not h['accepted_for_aggregation'] and h['diagnostic_readout_only'] for h in read)
    r=segment(feature(), [msg('Погода гарна')], [], PROFILE, read)
    assert r['primary_segment']=='D' and r['potential_post_alert_anchor_ids']==['m']


def test_boundary_windows_are_half_open_and_timezone_aware():
    rows=[msg('x',str(i),i) for i in [-11,-10,-1,0,9,10]]
    times=[r['message_time'] for r in rows]
    assert [r['message_id'] for r in window(rows,times,T-timedelta(minutes=10),T)]==['-10','-1']
    assert [r['message_id'] for r in window(rows,times,T,T+timedelta(minutes=10))]==['0','9']


def test_selection_reproducible_exclusions_fallback_and_unique_ids():
    rows=[dict(alert_id=f'{g}{n:03}', primary_segment=g,year=2023+n%4) for g in QUOTAS for n in range(15)]
    excluded={'A000','A001'}
    a,audit=select_residual(rows,excluded)
    b,_=select_residual(list(reversed(rows)),excluded)
    assert a==b and len(a)==40 and len({r['alert_id'] for r in a})==40
    assert not excluded & {r['alert_id'] for r in a}
    assert all(not r['fallback'] for r in a) and all(not r['shortfall'] for r in audit)
    limited=[r for r in rows if r['primary_segment']!='F']
    chosen,audit=select_residual(limited,set())
    assert len(chosen)==40 and audit[-1]['shortfall']==5 and sum(r['fallback'] for r in chosen)==5


@pytest.fixture(scope='module')
def real():
    data={k:pq.read_table(ROOT/p).to_pylist() for k,p in FILES.items()}
    def read(name):
        with (ROOT/OUTPUT/name).open(encoding='utf-8-sig',newline='') as stream:
            return list(csv.DictReader(stream))
    return data, read, json.loads((ROOT/OUTPUT/'residual_segment_summary.json').read_text())


def test_real_baseline_fail_closed(real):
    data,_,report=real
    assert baseline(data['v14'])==report['baseline']
    changed=copy.deepcopy(data['v14']);changed.pop()
    with pytest.raises(ValueError,match='BLOCKED'):
        baseline(changed)


def test_real_complete_all_residual_features_and_separate_coverage(real):
    data,read,report=real
    segments=read('residual_alert_segments.csv');gaps=read('no_linked_text_source_audit.csv')
    assert {r['alert_id'] for r in segments}=={f['alert_id'] for f in data['v14'] if f['classification_status']=='no_marker_with_text'}
    assert len(segments)==655 and len(gaps)==98
    assert not {r['alert_id'] for r in segments}&{r['alert_id'] for r in gaps}
    assert sum(report['primary_counts'].values())==655
    assert all(json.loads(r['diagnostic_segments']) and r['primary_segment'] in json.loads(r['diagnostic_segments']) for r in segments)
    assert sum(r['source_gap_type']=='only_empty_text_records' for r in gaps)==2


def test_real_sample_exclusions_quotas_and_risks(real):
    data,read,report=real
    rows=read('residual_review_alerts.csv');fm={f['alert_id']:f for f in data['v14']}
    assert len(rows)==50 and len({r['alert_id'] for r in rows})==50
    for r in rows:
        assert all(r[k]=='' for k in HUMAN_FIELDS)
        if r['sample_role']=='residual': assert fm[r['alert_id']]['classification_status']=='no_marker_with_text'
        else: assert not fm[r['alert_id']]['baseline_categories'] and fm[r['alert_id']]['unique_threat_category_count']
    assert not {r['alert_id'] for r in rows}&set(report['excluded_task0045_ids']+report['excluded_discussed_ids'])
    assert {g:sum(r['sample_role']=='residual' and r['primary_segment']==g for r in rows) for g in QUOTAS}==QUOTAS
    covered={risk for r in rows for risk in json.loads(r['safety_risk_flags'])}
    assert {'ballistic_high_speed','strategic_aviation','jet_uav','pre_alert','uncertain_geography','multiple_categories_uncertain_evidence'}<=covered


def test_real_complete_chronology_original_text_times_and_all_candidates(real):
    from collections import Counter
    data,read,_=real
    overview=read('residual_review_alerts.csv');ids={r['alert_id'] for r in overview}
    rows=read('residual_full_timeline.csv');pre=read('residual_pre_alert_evidence.csv')
    pairs={(r['alert_id'],r['message_id']) for r in data['links'] if r['alert_id'] in ids}
    assert len(rows)==len(pairs) and {(r['alert_id'],r['message_id']) for r in rows}==pairs
    mm={m['message_id']:m for m in data['messages'] if m['message_id'] is not None}
    for r in rows+pre:
        assert r['original_message_text']==(mm[r['message_id']]['text'] or '')
        assert r['timestamp_original']==mm[r['message_id']]['message_time_original']
        assert datetime.fromisoformat(r['timestamp_utc'])==mm[r['message_id']]['message_time']
    key=lambda h:(h['alert_id'],h['message_id'],h['rule_id'],h['match_start'],h['anchor_message_id'],h['evidence_role'])
    assert Counter(key(h) for h in data['markers'] if h['alert_id'] in ids)==Counter(key(h) for r in rows+pre for h in json.loads(r['candidate_records']))
    assert any(not json.loads(r['candidate_records']) for r in rows)


def test_workbook_review_fields_and_full_text_representation(real):
    _,read,report=real
    wb=load_workbook(ROOT/OUTPUT/'residual_review_50.xlsx')
    assert len(wb.sheetnames)==7 and wb['Human Review'].max_row==51
    assert wb['Human Review'].data_validations.count==2
    for ws in wb:
        assert ws.freeze_panes and ws.auto_filter.ref
        assert not any(c.data_type=='f' for row in ws for c in row)
    ws=wb['Human Review'];heads=[c.value for c in ws[1]]
    assert all(ws.cell(n,heads.index(k)+1).value is None for n in range(2,52) for k in HUMAN_FIELDS)
    timeline=read('residual_full_timeline.csv');ws=wb['Full Timeline'];heads=[c.value for c in ws[1]]
    assert ws.max_row==len(timeline)+1
    adjusted={(r['sheet'],r['row'],r['column']) for r in report['excel_cell_adjustments']}
    for n,r in enumerate(timeline,2):
        if ('Full Timeline',n,'original_message_text') not in adjusted:
            assert (ws.cell(n,heads.index('original_message_text')+1).value or '')==r['original_message_text']
    wb.close()


def test_frozen_sources_classifiers_and_prior_reports_unchanged(real):
    _,_,report=real
    current=snapshot(ROOT)
    # Protect every historical file while permitting explicitly authorized later outputs.
    assert {k:current[k] for k in report['protected_hashes']}==report['protected_hashes']
