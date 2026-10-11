"""Generalized reviewed statements, scoped negatives and bounded correspondence."""
from datetime import datetime, timedelta, timezone
from pathlib import Path
import copy

import pyarrow.parquet as pq
import pytest
import yaml

from kyiv_alerts.classification.rules import CATEGORIES
from kyiv_alerts.classification_v15.lexical import Profile
from kyiv_alerts.classification_v15.core import process

ROOT = Path(__file__).resolve().parents[1]
POLICY = yaml.safe_load((ROOT/'config/attribution_v1_5.yaml').read_text())
PROFILE = Profile(ROOT/'config/threat_dictionary_v1_3.yaml', ROOT/'config/kyiv_geography_v1_1.yaml',
                  yaml.safe_load((ROOT/'config/attribution_v1_4.yaml').read_text()), POLICY)
T = datetime(2026, 7, 1, 10, tzinfo=timezone.utc)
TEMPLATE = next(f for f in pq.read_table(ROOT/'data/processed/classification_v1_4/alert_threat_features.parquet').to_pylist()
                if f['classification_status'] == 'no_marker_with_text')


def synthetic(texts, earlier=(), extra_alerts=(), extra_links=(), baseline=(), policy=None):
    messages = []
    for i, (delta, text) in enumerate([*earlier, *texts]):
        messages.append(dict(message_id=str(i), text=text, message_time=T+timedelta(minutes=delta),
                             message_time_original='source '+str(i), message_type='message',
                             source_file='synthetic.html', source_message_index=i))
    linked = [m for m in messages if m['message_time'] >= T]
    f = copy.deepcopy(TEMPLATE)
    f.update(alert_id='synthetic', start_time=T, end_time=T+timedelta(hours=2), duration_minutes=120,
             telegram_message_count=len(linked), telegram_text_message_count=len(linked))
    links = [dict(alert_id='synthetic', message_id=m['message_id']) for m in linked]+list(extra_links)
    a = dict(alert_id='synthetic', start_time=T, end_time=f['end_time'], end_time_original=f['end_time'].isoformat())
    markers, features = process([f], messages, links, list(baseline), PROFILE, [a,*extra_alerts], policy or POLICY)
    cats = {c for c in CATEGORIES[:-1] if features[0][c] is True}
    return markers, features, cats, messages


@pytest.mark.parametrize('text,expected', [
    ('Поспали? Балістика з Курська', {'ballistic_high_speed'}),
    ('Загроза застосування балістичного озброєння', {'ballistic_high_speed'}),
    ('Повітряні сили повідомляють про загрозу застосування балістичного озброєння', {'ballistic_high_speed'}),
    ('Північ Києва — увага\nВорожий БпЛА', {'shahed_uav'}),
    ('БпЛА між Борисполем і Рембазою', {'shahed_uav'}),
    ('Реактивний з Житомирщини повернувся\nНа Глеваху', {'jet_uav'}),
])
def test_generalized_operational_statements(text, expected):
    markers, _, cats, _ = synthetic([(1,text)])
    assert cats == expected
    assert all(m['uncertainty_flag'] and m['uncertainty_reasons'] for m in markers if m['accepted_for_aggregation'])


@pytest.mark.parametrize('text', [
    'Вчора БпЛА над Києвом', 'Учора ворожий БпЛА над Києвом',
    'Київ: характеристики БпЛА та принцип роботи', 'Це вже не МіГ-31К',
    'МіГ-31К — без загроз', 'Стратегічна авіація без загроз',
    'БпЛА — загрози немає', 'БпЛА не підтверджено', 'Київ БпЛА?',
    'В новинах пишуть «ворожий БпЛА над Києвом»',
    'Одеса БпЛА курс на центр', 'БпЛА з Харкова на Полтаву',
    'Реактивний двигун повернувся до сервісу', 'Реактивна артилерія на Київ',
    'БпЛА не на Київ', 'Відбій: ворожий БпЛА над Києвом',
])
def test_negative_occurrences_remain_unaccepted(text):
    markers, _, cats, _ = synthetic([(1,text)])
    assert not cats and not any(m['accepted_for_aggregation'] for m in markers)


def test_one_family_negation_does_not_negate_another():
    markers, _, cats, _ = synthetic([(1,'Це вже не МіГ-31К, а ворожий БпЛА над Києвом')])
    assert cats == {'shahed_uav'}
    assert any(m['category']=='mig31' and not m['accepted_for_aggregation'] and m['scope_trace'] for m in markers)


def test_no_strikes_is_not_no_uav_and_not_a_ballistic_anchor():
    markers, _, cats, _ = synthetic([(1,'Ворожий БпЛА над Києвом. Ворожих ударів по столиці не було')])
    assert cats == {'shahed_uav'}
    assert any(t['kind']=='no_strikes_is_not_no_threat' and not t['blocks_occurrence'] for m in markers for t in m['scope_trace'])
    _, _, cats, _ = synthetic([(1,'Ворожих ударів по столиці не було')])
    assert not cats


def test_cross_message_route_change_supports_anchor_not_static_foreign_news():
    markers, _, cats, _ = synthetic([(1,'Х-59 в напрямку Чернігівщини'),(5,'Змінила курс на схід Київщини')])
    assert cats == {'cruise_missile'}
    assert any(m['evidence_role']=='context_reference' and m['accepted_for_aggregation'] for m in markers)
    _, _, cats, _ = synthetic([(1,'Одеса БпЛА'),(2,'Київ')])
    assert not cats


def test_pre_alert_explicit_cause_reference_can_cross_five_minute_gap_within_ten():
    markers, _, cats, _ = synthetic([(1,'Тривога в столиці через нього')], earlier=[(-4.5,'БпЛА на Бровари зі сходу')])
    assert cats == {'shahed_uav'}
    accepted = [m for m in markers if m['accepted_for_aggregation'] and m['time_relation']=='before_alert']
    assert accepted and 'explicit_pre_cause_reference_within_ten_minutes' in accepted[0]['attribution_rule_ids']


@pytest.mark.parametrize('delta,expected', [(-10,False),(-9,True),(-10.01,False)])
def test_explicit_pre_reference_bound_is_anchor_to_reference(delta, expected):
    _, _, cats, _ = synthetic([(0.5,'Тривога в столиці через нього')], earlier=[(delta,'БпЛА на Бровари')])
    assert bool(cats) is expected


def test_pre_warning_alone_cannot_add_category():
    _, _, cats, _ = synthetic([(1,'Погода гарна')], earlier=[(-1,'Київ — загроза балістики')])
    assert not cats
    _, _, cats, _ = synthetic([(1,'Ворожий БпЛА над Києвом')], earlier=[(-1,'Загроза застосування балістичного озброєння')])
    assert cats == {'shahed_uav'}


def test_pre_alert_previous_membership_cannot_transfer():
    previous = dict(alert_id='previous', start_time=T-timedelta(minutes=20), end_time=T-timedelta(minutes=1), end_time_original=(T-timedelta(minutes=1)).isoformat())
    markers, _, cats, _ = synthetic([(1,'Тривога в столиці через нього')], earlier=[(-2,'БпЛА на Бровари')],
                                    extra_alerts=[previous], extra_links=[dict(alert_id='previous',message_id='0')])
    assert not cats
    assert any(m['previous_alert_ids'] or m['intervening_boundary_alert_ids'] for m in markers)


def test_direction_reversal_and_lost_observation_do_not_erase_prior_threat():
    markers, _, cats, _ = synthetic([(1,'Реактивний з Житомирщини повернувся\nНа Глеваху'),
                                     (3,'Летовище'),(3.1,'Знижується'),(4,'Не фіксується'),(5,'Наступний від Славутича')])
    assert cats == {'jet_uav'}
    assert not any(m['message_id']=='4' and m['accepted_for_aggregation'] for m in markers)


def test_multiple_simultaneous_families_leave_pronoun_unresolved():
    markers, _, cats, _ = synthetic([(1,'БпЛА на Бровари і Х-59 на Чернігівщину'),
                                     (2,'Тривога в столиці через нього')])
    assert cats == {'shahed_uav', 'cruise_missile'}
    assert all(not m['accepted_for_aggregation'] for m in markers if m['evidence_role']=='context_reference')


def test_late_or_unrelated_followup_does_not_bridge_episode():
    _, _, cats, _ = synthetic([(1,'Х-59 в напрямку Чернігівщини'),(7,'Змінила курс на схід Київщини')])
    assert not cats
    _, _, cats, _ = synthetic([(1,'Х-59 в напрямку Чернігівщини'),(2,'Наступна на Київ')])
    assert not cats


def test_specific_trace_spans_preserve_source_expression():
    markers, _, _, messages = synthetic([(1,'Це вже не МіГ-31К. Загрози немає')])
    mm={m['message_id']:m for m in messages}
    traces=[(m,t) for m in markers for t in m['scope_trace']]
    assert traces
    for m,t in traces:
        assert mm[m['message_id']]['text'][t['match_start']:t['match_end']]==t['matched_expression']


def test_scope_between_families_without_separate_sentence():
    _,_,cats,_=synthetic([(1,'МіГ-31К без загроз, ворожий БпЛА над Києвом')])
    assert cats=={'shahed_uav'}


def test_negated_pending_family_cannot_resume_via_locality():
    markers,_,cats,_=synthetic([(1,'Х-59 в напрямку Чернігівщини'),
                              (2,'Х-59 — загрози немає'),(3,'Змінила курс на схід Київщини')])
    assert not cats


def test_no_threat_followup_cannot_corroborate_pre_warning():
    _,_,cats,_=synthetic([(1,'Київ без загроз')],earlier=[(-2,'БпЛА на Бровари')])
    assert not cats


def test_absence_of_communications_does_not_negate_uav():
    _,_,cats,_=synthetic([(1,"БпЛА над Києвом, зв'язку немає")])
    assert cats=={'shahed_uav'}


@pytest.fixture(scope='module')
def real_v15():
    import json
    from kyiv_alerts.classification_v15.__main__ import FILES
    data={k:pq.read_table(ROOT/p).to_pylist() for k,p in FILES.items()}
    markers=pq.read_table(ROOT/'data/processed/classification_v1_5/message_threat_markers.parquet').to_pylist()
    features=pq.read_table(ROOT/'data/processed/classification_v1_5/alert_threat_features.parquet').to_pylist()
    report=json.loads((ROOT/'reports/classification_v1_5/classification_comparison.json').read_text())
    return data,markers,features,report


def test_real_source_reconciliation_and_frozen_baselines(real_v15):
    from kyiv_alerts.classification_v15.__main__ import validate,snapshot,rule_hashes
    data,markers,features,report=real_v15
    validate(markers,features,data)
    current=snapshot(ROOT)
    assert {k:current[k] for k in report['protected_hashes']}==report['protected_hashes']
    assert rule_hashes(ROOT)==report['rule_hashes']
    assert len(features)==2594 and report['v1_5']['eligible_alerts']==1803


def test_real_gains_losses_categories_and_calendar_totals_reconcile(real_v15):
    from kyiv_alerts.classification.consolidation import summary,compare_versions
    data,markers,features,report=real_v15
    assert summary(features)==report['v1_5']
    comparison,_=compare_versions(data['v14'],features)
    assert comparison==report['comparison_v14_v15']
    assert report['net_new']==comparison['gained_count']-comparison['lost_count']
    assert sum(r['v15_classified'] for r in report['monthly'])==report['v1_5']['classified_alerts']
    assert sum(r['v15_classified'] for r in report['yearly'])==report['v1_5']['classified_alerts']
    for c in CATEGORIES[:-1]:
        assert len({m['alert_id'] for m in markers if m['accepted_for_aggregation'] and m['category']==c})==sum(f[c] is True for f in features)


def test_reviewed_targets_discrepancies_and_provisional_labels_remain_explicit(real_v15):
    _,_,_,report=real_v15
    rows={r['alert_id']:r for r in report['regression_cases']}
    for aid in ['833','1312','1500','2589','806']:
        assert rows[aid]['regression_status']=='CONFIRMED_LABEL_REPRODUCED'
    assert rows['1096']['expected_categories']==['shahed_uav','ballistic_high_speed']
    assert rows['1096']['missing_expected_categories']==['ballistic_high_speed']
    assert rows['1096']['regression_status']=='CONFIRMED_LABEL_DISCREPANCY'
    assert rows['1416']['label_status']=='provisional_pending_formal_approval'
    assert rows['2581']['existing_research_lead_decision']=='retain_discrepancy_for_review'
    assert sum(r['task0045_sample'] for r in rows.values())==25


def test_holdout_reservation_and_audit_are_disjoint_and_unlabeled(real_v15):
    import csv
    from openpyxl import load_workbook
    from kyiv_alerts.residual_review.core import HUMAN_FIELDS
    _,_,_,report=real_v15
    holdout=set(report['holdout']['reserved_alert_ids'])
    assert len(holdout)==43 and not holdout&set(report['holdout']['development_ids'])
    audit=set(report['audit']['selected_ids'])
    assert len(audit)==20 and not audit&holdout
    assert audit<=set(report['comparison_v14_v15']['gained_alert_ids'])
    assert not audit&{r['alert_id'] for r in report['regression_cases']}
    wb=load_workbook(ROOT/'reports/classification_v1_5/false_positive_audit_20.xlsx')
    ws=wb['Human Review'];heads=[c.value for c in ws[1]]
    assert all(ws.cell(n,heads.index(k)+1).value is None for n in range(2,22) for k in HUMAN_FIELDS)
    with (ROOT/'reports/classification_v1_5/audit_full_timeline.csv').open(encoding='utf-8-sig',newline='') as stream:
        timeline=list(csv.DictReader(stream))
    assert len(timeline)==report['audit']['timeline_messages']==wb['Full Timeline'].max_row-1
    wb.close()


def test_pre_only_requires_inside_source_correspondence_and_uncertainty(real_v15):
    data,markers,features,report=real_v15
    pairs={(l['alert_id'],l['message_id']) for l in data['links']}
    for m in markers:
        if m['accepted_for_aggregation'] and m['time_relation']=='before_alert':
            assert any((m['alert_id'],mid) in pairs for mid in m['support_message_ids'])
            assert not m['previous_alert_ids'] and not m['intervening_boundary_alert_ids'] and not m['unresolved_envelope_alert_ids']
    gained=set(report['comparison_v14_v15']['gained_alert_ids'])
    uncertain={m['alert_id'] for m in markers if m['accepted_for_aggregation'] and m['uncertainty_flag'] and m['alert_id'] in gained}
    assert len(uncertain)==report['new_uncertain_classifications']
    # Development IDs are restricted to audit fixtures, never interpretation code.
    for filename in ['core.py','lexical.py']:
        code=(ROOT/'src/kyiv_alerts/classification_v15'/filename).read_text()
        assert all("'"+aid+"'" not in code for aid in ['833','1312','1500','2589','806','1096','1416'])


def test_baseline_failure_writes_diagnostic_and_never_runs_classifier(real_v15,tmp_path,monkeypatch):
    import json
    from types import SimpleNamespace
    from kyiv_alerts.classification_v15 import __main__ as cli
    data,_,_,_=real_v15
    changed=copy.deepcopy(data);changed['v14'].pop()
    bypath={v:k for k,v in cli.FILES.items()}
    monkeypatch.setattr(cli,'snapshot',lambda root:{})
    monkeypatch.setattr(cli,'rule_hashes',lambda root:{})
    monkeypatch.setattr(cli.pq,'read_table',lambda path:SimpleNamespace(to_pylist=lambda:changed[bypath[str(path.relative_to(tmp_path))]]))
    monkeypatch.setattr(cli,'process',lambda *args,**kwargs:pytest.fail('Classifier ran after mismatched baseline'))
    with pytest.raises(ValueError,match='BLOCKED'):
        cli.run(tmp_path)
    abort=json.loads((tmp_path/'reports/classification_v1_5/baseline_abort.json').read_text())
    assert abort['status']=='BLOCKED' and not (tmp_path/'data/processed/classification_v1_5').exists()
