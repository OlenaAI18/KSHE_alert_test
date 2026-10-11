"""Trace/provenance contracts, deterministic sampling and native Excel usability."""
from pathlib import Path
from collections import Counter
from datetime import datetime,timedelta,timezone
import copy,json,csv
import pytest
from openpyxl import load_workbook,Workbook
from kyiv_alerts.controlled_review.core import select,reconstruct,SEED,STATUS,association_key
from kyiv_alerts.controlled_review.__main__ import FILES,validate_inputs,validate_review,protected_snapshot
from kyiv_alerts.controlled_review.workbook import write_sheet,dropdown
import pyarrow.parquet as pq

ROOT=Path(__file__).resolve().parents[1]
@pytest.fixture(scope='module')
def data():
    return {k:pq.read_table(ROOT/p).to_pylist() for k,p in FILES.items() if p.endswith('.parquet')}

@pytest.fixture(scope='module')
def built(data):
    selection,strata=select(data['features'],data['markers'],data['inside'])
    overview,timeline,trace,pre=reconstruct(selection,data['features'],data['messages'],data['links'],data['markers'],data['inside'],data['pre'])
    return selection,strata,overview,timeline,trace,pre


def test_exact_controlled_sample(built,data):
    selected,strata,*_=built
    assert len(selected)==len({r['alert_id'] for r in selected})==25
    assert Counter(r['sample_role'] for r in selected)=={'unclassified':20,'positive_control':5}
    assert Counter(r['selection_stratum'] for r in selected)=={'october_2026':8,'august_september_2026':5,'high_volume_2024_2025':4,'low_volume_2023_2025':3,'positive_control':5}
    assert not any(r['unavailable'] for r in strata)
    assert {r['alert_id'] for r in selected}<={r['alert_id'] for r in data['alerts']}


def test_stratum_bounds_and_text(built,data):
    fm={r['alert_id']:r for r in data['features']}
    for selected in built[0]:
        f=fm[selected['alert_id']];s=selected['selection_stratum']
        if selected['sample_role']=='unclassified':
            assert f['classification_status']=='no_marker_with_text' and f['telegram_text_message_count']>0
        if s=='high_volume_2024_2025':assert f['telegram_text_message_count']>=20 and f['start_time'].year in {2024,2025}
        if s=='low_volume_2023_2025':assert f['telegram_text_message_count']<=3 and f['start_time'].year in {2023,2024,2025}


def test_required_diagnostic_signals_present(built):
    signals={tag for r in built[0] for tag in r['diagnostic_signals']}
    assert {'no_lexical_candidates','rejected_lexical_candidates','task0044_provisional_inside','ambiguous_geography','uncertain_context','explicit_named_threat_lexical_hit'}<=signals


def test_selection_seed_and_reordered_input(data):
    a=select(data['features'],data['markers'],data['inside'])
    assert a==select(data['features'],data['markers'],data['inside'])
    assert a==select(list(reversed(data['features'])),list(reversed(data['markers'])),list(reversed(data['inside'])))


def test_unavailable_strata_reported_not_replaced(data):
    selected,strata=select([],[],[])
    assert selected==[] and sum(r['unavailable'] for r in strata)==25
    assert [r['selected'] for r in strata]==[0]*5


def test_all_linked_messages_not_only_hits(built,data):
    selected,_,_,timeline,_,_=built;ids={r['alert_id'] for r in selected}
    expected={(r['alert_id'],r['message_id']) for r in data['links'] if r['alert_id'] in ids}
    assert len(timeline)==len(expected) and {(r['alert_id'],r['message_id']) for r in timeline}==expected
    assert any(not r['dictionary_hit'] for r in timeline)
    assert any(not r['original_message_text'] for r in timeline)


def test_every_candidate_and_accepted_marker_retained(built,data):
    selected,_,_,_,trace,_=built;ids={r['alert_id'] for r in selected}
    expected=[r for r in data['markers'] if r['alert_id'] in ids]
    key=lambda r:(r['alert_id'],r['message_id'],r['rule_id'],r['match_start'],r['match_end'])
    assert Counter(map(key,expected))==Counter(map(key,trace))
    assert sum(r['accepted_for_aggregation'] for r in expected)==sum(r['final_acceptance_status']=='accepted' for r in trace)


def test_all_source_text_original_times_provenance(built,data):
    mm={r['message_id']:r for r in data['messages'] if r['message_id'] is not None}
    for r in built[3]:
        m=mm[r['message_id']]
        assert (r['original_message_text'],r['timestamp'],r['timestamp_original'],r['source_file'],r['source_message_index'])==(m['text'],m['message_time'],m['message_time_original'],m['source_file'],m['source_message_index'])
    for r in built[4]:
        assert r['original_text'][r['match_start']:r['match_end']]==r['matched_expression']


def test_chronology_and_relative_time(built,data):
    fm={r['alert_id']:r for r in data['features']};times={}
    for r in built[3]:
        aid=r['alert_id']
        if aid in times:assert r['timestamp']>=times[aid]
        times[aid]=r['timestamp']
        assert r['minutes_from_alert_start']==(r['timestamp']-fm[aid]['start_time']).total_seconds()/60


def test_no_human_labels_or_auto_root_causes(built):
    for rows in [built[2],built[3],built[4]]:
        for r in rows:assert all(v is None for k,v in r.items() if k.startswith('human_') or k=='review_notes')
    assert all(r['proposed_diagnostic_group'].startswith('Review questions only:') for r in built[2] if r['sample_role']=='unclassified')


def test_provisional_associations_remain_provisional(built):
    for r in built[4]:
        for a in r['task0044_provisional_associations']:
            assert a['confidence_or_review_status']=='PROVISIONAL — HUMAN VALIDATION REQUIRED'
            assert a['human_decision'] is None
        if r['task0044_provisional_associations']:assert 'occurrence spans unavailable' in r['task0044_association_alignment']


def test_missing_instrumentation_is_not_fabricated(built):
    for r in built[4]:
        assert r['context_rule']['legacy_pattern_ids']=='not_recorded'
        assert r['geography_rule']['legacy_regex_rule_ids']=='not_recorded'
        assert r['context_rule']['approved_override_rule_id']==r['context_rule_id_recorded']
        assert r['geography_rule']['relevance_rule_id']==r['relevance_rule_id_recorded']


def test_read_only_and_foreign_key_validation(built,data):
    before=copy.deepcopy(built[0]);validate_inputs(data)
    validate_review(built[0],built[2],built[3],built[4],data)
    assert built[0]==before


def test_missing_message_is_detected(built,data):
    with pytest.raises(ValueError,match='Timeline omitted'):
        validate_review(built[0],built[2],built[3][1:],built[4],data)


def test_missing_accepted_occurrence_detected(built,data):
    idx=next(i for i,r in enumerate(built[4]) if r['final_acceptance_status']=='accepted')
    trace=built[4][:idx]+built[4][idx+1:]
    with pytest.raises(ValueError,match='Candidate occurrence omitted'):
        validate_review(built[0],built[2],built[3],trace,data)


def test_invented_human_decision_detected(built,data):
    ov=copy.deepcopy(built[2]);ov[0]['human_decision']='confirm_context'
    with pytest.raises(ValueError,match='Human label invented'):
        validate_review(built[0],ov,built[3],built[4],data)


def test_workbook_opens_and_usability(built):
    wb=load_workbook(ROOT/'reports/controlled_review/controlled_alert_review.xlsx')
    assert wb.sheetnames==['README','Alert Overview','Full Message Timeline','Classification Trace','Diagnostic Summary','Pre-Alert Evidence']
    assert wb['Alert Overview'].max_row==len(built[2])+1
    assert wb['Full Message Timeline'].max_row==len(built[3])+1
    assert wb['Classification Trace'].max_row==len(built[4])+1
    for ws in wb:
        assert ws.freeze_panes=='C2' and ws.auto_filter.ref
        assert ws['A1'].font.bold
    ov=wb['Alert Overview'];assert len(ov.data_validations.dataValidation)==3
    assert ov['A2'].hyperlink and ov['A2'].hyperlink.target.startswith("#'Full Message Timeline'")
    for ws in [ov,wb['Full Message Timeline'],wb['Classification Trace']]:
        headers=[c.value for c in ws[1]]
        for col,k in enumerate(headers,1):
            if k.startswith('human_') or k=='review_notes':assert all(ws.cell(row,col).value is None for row in range(2,ws.max_row+1))
    wb.close()


def test_workbook_source_text_and_dates(built):
    wb=load_workbook(ROOT/'reports/controlled_review/controlled_alert_review.xlsx');ws=wb['Full Message Timeline']
    headers=[c.value for c in ws[1]]
    for i,r in enumerate(built[3],2):
        cell=ws.cell(i,headers.index('original_message_text')+1)
        assert (cell.value or '')==(r['original_message_text'] or '')
        assert cell.alignment.wrap_text
        assert ws.cell(i,headers.index('timestamp')+1).number_format=='yyyy-mm-dd hh:mm:ss'
    assert not wb._external_links
    assert not any(c.data_type=='f' for ws in wb for row in ws for c in row)
    wb.close()


def test_excel_long_text_fallback_and_formula_literal(tmp_path):
    wb=Workbook();wb.remove(wb.active);adjust=[]
    text='='+'x'*40000
    ws=write_sheet(wb,'Example',[dict(alert_id='a',message_id='m',original_text=text,human_notes=None)],adjustments=adjust)
    assert len(ws['C2'].value)==32767 and ws['C2'].data_type=='s'
    assert 'companion CSV' in ws['C2'].value and len(adjust)==1
    ws2=write_sheet(wb,'Formula',[{'original_text':'=1+1'}]);assert ws2['A2'].data_type=='s'
    path=tmp_path/'test.xlsx';wb.save(path);loaded=load_workbook(path)
    assert loaded['Formula']['A2'].value=='=1+1' and loaded['Formula']['A2'].data_type=='s'
    loaded.close();assert text=='='+'x'*40000


def test_excel_xml_control_chars_preserved_in_source():
    wb=Workbook();wb.remove(wb.active);adjust=[];row={'original_text':'hello\x01world'}
    ws=write_sheet(wb,'Example',[row],adjustments=adjust)
    assert ws['A2'].value=='hello\\u0001world' and row['original_text']=='hello\x01world'
    assert len(adjust)==1


def test_real_report_counts_and_protected_hashes(built):
    r=json.loads((ROOT/'reports/controlled_review/controlled_review_summary.json').read_text())
    assert r['review_status']==STATUS and r['selected_alert_ids']==[s['alert_id'] for s in built[0]]
    assert r['total_linked_messages']==len(built[3]) and r['candidate_occurrences']==len(built[4])
    assert r['accepted_occurrences']+r['rejected_occurrences']==r['candidate_occurrences']
    assert r['messages_without_dictionary_hit']+r['messages_with_dictionary_hit']==r['total_linked_messages']
    current=protected_snapshot(ROOT)
    assert r['protected_hashes']=={k:current[k] for k in r['protected_hashes']}
    assert r['human_reviewed_alerts']==0 and r['population_precision'] is None and r['population_recall'] is None
