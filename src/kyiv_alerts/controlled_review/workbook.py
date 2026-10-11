"""Openpyxl review workbook with blank validated human fields and full CSV fallback."""
from datetime import datetime
import json,math
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from kyiv_alerts.ingestion.timeutils import KYIV
from .core import STATUS,GROUPS

DATE_FORMAT='yyyy-mm-dd hh:mm:ss'
BLUE='17365D';WHITE='FFFFFF';ACCEPT='E2F0D9';REJECT='FCE4D6';REVIEW='FFF2CC'


def value(v):
    if isinstance(v,(list,dict)):return json.dumps(v,ensure_ascii=False,sort_keys=True,default=lambda x:x.isoformat() if hasattr(x,'isoformat') else str(x))
    if isinstance(v,datetime):return v.astimezone(KYIV).replace(tzinfo=None)
    return v


def write_sheet(wb,name,rows,columns=None,adjustments=None):
    ws=wb.create_sheet(name);columns=columns or list(rows[0]) if rows else columns or ['information']
    ws.append(columns);ws.freeze_panes='C2';ws.auto_filter.ref=f'A1:{ws.cell(1,len(columns)).column_letter}{len(rows)+1}'
    changed=[]
    for column,k in enumerate(columns,1):
        c=ws.cell(1,column);c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(color=WHITE,bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
        width=82 if k in {'original_text','original_message_text','evidence_text'} else 58 if k in {'evidence_context','task0044_evidence','task0044_provisional_associations'} else 30 if k in {'timestamp_original','source_file','review_notes','human_notes'} else 24
        ws.column_dimensions[c.column_letter].width=width
    ws.row_dimensions[1].height=42
    for ri,row in enumerate(rows,2):
        longest=0
        for ci,k in enumerate(columns,1):
            v=value(row.get(k));cell=ws.cell(ri,ci)
            if isinstance(v,str):
                cleaned=ILLEGAL_CHARACTERS_RE.sub(lambda m:f'\\u{ord(m.group()):04x}',v)
                if cleaned!=v:changed.append(dict(sheet=name,row=ri,column=k,alert_id=row.get('alert_id'),message_id=row.get('message_id'),reason='Excel XML control-character escape; CSV retains original'))
                if len(cleaned)>32767:
                    suffix='\n[Excel cell limit; full original in companion CSV, same alert/message ID.]'
                    cleaned=cleaned[:32767-len(suffix)]+suffix;changed.append(dict(sheet=name,row=ri,column=k,alert_id=row.get('alert_id'),message_id=row.get('message_id'),reason='Excel 32767-character limit; full original in companion CSV'))
                cell.value=cleaned;cell.data_type='s' # Source strings beginning = are text, never executable formulas.
                if k in {'original_text','original_message_text','evidence_text','evidence_context'}:longest=max(longest,len(cleaned))
            else:cell.value=v
            if isinstance(v,datetime):cell.number_format=DATE_FORMAT
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if k.startswith('human_') or k in {'review_notes'}:cell.fill=PatternFill('solid',fgColor=REVIEW)
            if k in {'final_acceptance_status','accepted_rejected','current_status','task0044_candidate_status','task0044_status'}:
                fill=ACCEPT if v=='accepted' or isinstance(row.get(k),list) and row[k] and set(row[k])=={'accepted'} or k=='current_status' and v in {'single_type','multiple_types','only_other'} else REJECT if v=='rejected' or isinstance(row.get(k),list) and row[k] and set(row[k])=={'rejected'} else REVIEW
                cell.fill=PatternFill('solid',fgColor=fill)
        ws.row_dimensions[ri].height=min(260,max(46,math.ceil(longest/110)*15))
    if adjustments is not None:adjustments.extend(changed)
    return ws


def dropdown(ws,field,choices):
    headers=[c.value for c in ws[1]]
    if field not in headers or ws.max_row<2:return
    col=ws.cell(1,headers.index(field)+1).column_letter
    dv=DataValidation(type='list',formula1='"'+','.join(choices)+'"',allow_blank=True)
    dv.errorTitle='Choose a review value';dv.error='Select a listed value; details and multiple categories go in review notes.'
    dv.showErrorMessage=True;dv.showInputMessage=True;dv.promptTitle='Human review only';dv.prompt='Leave blank until a reviewer examines the full source timeline.'
    ws.add_data_validation(dv);dv.add(f'{col}2:{col}{ws.max_row}')


def build(path,overview,timeline,trace,pre,summary):
    wb=Workbook();wb.remove(wb.active);adjust=[]
    instructions=[
      {'topic':'Статус','explanation':STATUS},
      {'topic':'Мета','explanation':'З’ясувати, де методологія втрачає інформацію. Це діагностика, не новий класифікатор і не підтвердження помилок.'},
      {'topic':'Вибірка','explanation':f"Цільова діагностична вибірка, seed={summary['seed']}. 8 жовтень 2026; 5 серпень–вересень 2026; 4 за 2024–2025 з ≥20 текстовими повідомленнями; 3 за 2023–2025 з ≤3; 5 позитивних контрольних. Дані жовтня закінчуються 4–5 жовтня."},
      {'topic':'Алгоритм відбору','explanation':'Сортування ID, перемішування фіксованим seed, добір нових діагностичних сигналів. Позитивні контрольні: різні категорії, явна причина, коротші інтервали за кількістю повідомлень. Порожні страти не замінюються іншими.'},
      {'topic':'Поточна класифікація','explanation':'v1.3: accepted означає проходження чинних словникових правил, не незалежне підтвердження типу зброї. shahed_uav — широка аналітична група БПЛА, generic_uav не визначає модель.'},
      {'topic':'Кандидати 004.4','explanation':'Усі записи лишаються PROVISIONAL — HUMAN VALIDATION REQUIRED. Рівні supported не означають затвердження. Pre-Alert Evidence — передтривожні записи, а не повідомлення всередині тривоги.'},
      {'topic':'Як перевіряти','explanation':'Від Alert Overview перейдіть до Full Message Timeline. Прочитайте ВСІ повідомлення по порядку. Перевірте явні назви загроз та попередні/наступні уточнення. Зіставте Classification Trace з оригіналом. Заповнюйте жовті поля тільки після перевірки.'},
      {'topic':'Час','explanation':'Дати Excel відображено в Europe/Kyiv без timezone через обмеження Excel. timestamp_original збережено буквально; CSV містить нормалізовані UTC/offset timestamp. H2 не змінено.'},
      {'topic':'Root causes A–H','explanation':'Це гіпотези, не істинні мітки. Proposed diagnostic group = питання для перевірки. Відхилення географії/контексту може бути правильним.'},
      {'topic':'Human decision','explanation':'confirm_context / reject_context / insufficient_information / needs_discussion / positive_control_confirmed. Початково всі поля порожні.'},
      {'topic':'Human threat category','explanation':'Оберіть категорію, multiple_categories, insufficient_information або no_operational_threat. Для кількох категорій запишіть їх у Review notes. Не створюйте нових категорій.'},
      {'topic':'Human interpretation','explanation':'operational_threat / contextual_reference / non_operational / uncertain. Ці значення заповнює лише людина.'},
      {'topic':'Missing instrumentation','explanation':'not_recorded означає відсутній запис внутрішнього рішення. Існуючі override IDs та назви географії збережені; legacy regex pattern IDs та конкретний suppressor не були записані. У 004.4 немає match spans для однозначного вирівнювання повторних входжень.'},
      {'topic':'Повний текст','explanation':'Текст збережено в companion CSV. Якщо клітинка перевищує Excel 32767 символів або має заборонені XML control characters, зміни лише у Excel явно перелічено в JSON; CSV зберігає оригінал. Формули, макроси й зовнішні зв’язки не створюються.'},
      {'topic':'Статистичні межі','explanation':'Не оцінюйте population precision/recall чи coverage improvement з цієї цільової вибірки. Висновки потребують ручної перевірки.'},
    ]+[{'topic':k,'explanation':v} for k,v in GROUPS.items()]
    readme=write_sheet(wb,'README',instructions);readme.column_dimensions['A'].width=29;readme.column_dimensions['B'].width=125
    for ri in range(2,readme.max_row+1):readme.row_dimensions[ri].height=65
    ov=write_sheet(wb,'Alert Overview',overview,adjustments=adjust)
    tl=write_sheet(wb,'Full Message Timeline',timeline,adjustments=adjust)
    tr=write_sheet(wb,'Classification Trace',trace,adjustments=adjust)
    summaryrows=[dict(metric=k,value=v) for k,v in summary.items() if k not in {'input_hashes','protected_hashes','elapsed_seconds'}]
    sm=write_sheet(wb,'Diagnostic Summary',summaryrows,adjustments=adjust);sm.column_dimensions['A'].width=42;sm.column_dimensions['B'].width=120
    pr=write_sheet(wb,'Pre-Alert Evidence',pre,adjustments=adjust)
    first={}
    for i,row in enumerate(timeline,2):first.setdefault(row['alert_id'],i)
    for i,row in enumerate(overview,2):
        cell=ov.cell(i,1);cell.hyperlink=f"#'Full Message Timeline'!A{first[row['alert_id']]}";cell.font=Font(color='0563C1',underline='single')
    for ws in [tl,tr,pr]:
        for i in range(2,ws.max_row+1):
            aid=ws.cell(i,1).value
            index=next((j for j,r in enumerate(overview,2) if r['alert_id']==aid),None)
            if index:ws.cell(i,1).hyperlink=f"#'Alert Overview'!A{index}"
    dropdown(ov,'human_confirmed_threat_category',['shahed_uav','jet_uav','ballistic_high_speed','cruise_missile','mig31','strategic_aviation','other','multiple_categories','insufficient_information','no_operational_threat'])
    dropdown(ov,'human_confirmed_primary_failure_reason',list(GROUPS)+['no_failure','not_applicable'])
    dropdown(ov,'human_decision',['confirm_context','reject_context','insufficient_information','needs_discussion','positive_control_confirmed'])
    for ws in [tl,tr]:dropdown(ws,'human_interpretation',['operational_threat','contextual_reference','non_operational','uncertain'])
    wb.properties.title='Controlled Alert Review — Task 004.5';wb.properties.description=STATUS
    wb.save(path)
    check=load_workbook(path)
    if check.sheetnames!=['README','Alert Overview','Full Message Timeline','Classification Trace','Diagnostic Summary','Pre-Alert Evidence']:raise ValueError('Workbook structure failed')
    check.close()
    return adjust
