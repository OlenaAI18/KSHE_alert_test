"""Read frozen artifacts, export full evidence, validate and stop for human review."""
import argparse
from bisect import bisect_left
from collections import Counter, defaultdict
from datetime import timedelta
from pathlib import Path
import csv
import json
import time

import pyarrow.parquet as pq
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .core import (SEED, STATUS, QUOTAS, RULES, PRIMARY_PRIORITY, GAP_PROBES,
                   HUMAN_FIELDS, blank_human, segment, message_signals,
                   select_residual, select_safety)
from kyiv_alerts.controlled_review.__main__ import digest
from kyiv_alerts.controlled_review.core import order
from kyiv_alerts.controlled_review.workbook import write_sheet, dropdown
from kyiv_alerts.classification.coverage_audit import write_audit_csv
from kyiv_alerts.classification_v14.lexical import LexicalProfile
from kyiv_alerts.classification_v14.core import operational
from kyiv_alerts.integration.__main__ import validate_h2, validate_links
from kyiv_alerts.attribution.core import contamination
from kyiv_alerts.ingestion.timeutils import KYIV

OUTPUT = Path('reports/classification_v1_4/residual_review')
FILES = {
    'alerts': 'data/interim/alerts.parquet',
    'messages': 'data/interim/telegram_messages.parquet',
    'links': 'data/processed/alert_message_links.parquet',
    'coverage': 'data/processed/alert_coverage.parquet',
    'v13': 'data/processed/classification_v1_3/alert_threat_features.parquet',
    'v14': 'data/processed/classification_v1_4/alert_threat_features.parquet',
    'markers': 'data/processed/classification_v1_4/message_threat_markers.parquet',
}
EXPECTED = dict(official=2594, eligible=1803, classified=1050, text_unclassified=655,
                no_linked_text=98, outside_coverage=790, dst_unresolved=1)


def snapshot(root):
    folders = ['data', 'config', 'reports', 'src/kyiv_alerts/classification',
               'src/kyiv_alerts/classification_v14', 'src/kyiv_alerts/integration',
               'src/kyiv_alerts/ingestion']
    return {str(p.relative_to(root)): digest(p) for folder in folders
            for p in sorted((root/folder).rglob('*'))
            if p.is_file() and '__pycache__' not in p.parts and not p.is_relative_to(root/OUTPUT)}


def baseline(features):
    counts = Counter(f['classification_status'] for f in features)
    classified = sum(counts[k] for k in ['single_type', 'multiple_types', 'only_other'])
    actual = dict(official=len(features), eligible=classified+counts['no_marker_with_text']+counts['no_linked_text'],
                  classified=classified, text_unclassified=counts['no_marker_with_text'],
                  no_linked_text=counts['no_linked_text'], outside_coverage=counts['outside_telegram_coverage'],
                  dst_unresolved=counts['unresolved'])
    if actual != EXPECTED:
        raise ValueError('BLOCKED: baseline differs; no review generated: '+json.dumps(actual))
    return actual


def window(timed, times, lower, upper):
    return timed[bisect_left(times, lower):bisect_left(times, upper)]


def diagnostic_readouts(messages, dictionary, cache):
    result = []
    for m in messages:
        mid = m['message_id']
        if mid not in cache:
            cache[mid] = []
            for h in dictionary.classify(m['text'] or ''):
                ok, rule = operational(h, m)
                cache[mid].append(dict(h, message_id=mid, operational_status='operational' if ok else 'non_operational_or_unresolved',
                                       operational_rule_id=rule, diagnostic_readout_only=True,
                                       accepted_for_aggregation=False))
        result.extend(cache[mid])
    return result


def source_coverage_audit(features, seq, timed, times):
    rows = []
    for f in features:
        if f['classification_status'] != 'no_linked_text':
            continue
        a = bisect_left(times, f['start_time'])
        b = bisect_left(times, f['end_time'])
        previous = timed[a-1] if a else None
        following = timed[b] if b < len(timed) else None
        inside = seq[f['alert_id']]
        rows.append(dict(alert_id=f['alert_id'], start_local=f['start_time'].astimezone(KYIV).isoformat(),
                         end_local=f['end_time'].astimezone(KYIV).isoformat(), duration_minutes=f['duration_minutes'],
                         coverage_status=f['coverage_status'], source_gap_type='only_empty_text_records' if inside else 'no_linked_records',
                         linked_message_count=len(inside), linked_text_count=0,
                         timestamped_source_records_in_interval=b-a,
                         previous_message_id=previous['message_id'] if previous else None,
                         previous_time_local=previous['message_time'].astimezone(KYIV).isoformat() if previous else None,
                         minutes_from_previous_to_start=(f['start_time']-previous['message_time']).total_seconds()/60 if previous else None,
                         next_message_id=following['message_id'] if following else None,
                         next_time_local=following['message_time'].astimezone(KYIV).isoformat() if following else None,
                         minutes_from_end_to_next=(following['message_time']-f['end_time']).total_seconds()/60 if following else None,
                         surrounding_publication_gap_minutes=(following['message_time']-previous['message_time']).total_seconds()/60 if previous and following else None,
                         interpretation='Source text unavailable; not evidence of no threat; no A-F diagnosis; media/OCR not read'))
    return rows


def reconstruct(selected, data, seq, bymarker, timed, times, dictionary, cache):
    fm = {f['alert_id']: f for f in data['v14']}
    old = {f['alert_id']: f for f in data['v13']}
    official = {a['alert_id']: a for a in data['alerts']}
    mm = {m['message_id']: m for m in data['messages'] if m['message_id'] is not None}
    members = defaultdict(set)
    for l in data['links']:
        members[l['message_id']].add(l['alert_id'])
    overview, timeline, pre, post = [], [], [], []
    for s in selected:
        aid = s['alert_id']; f = fm[aid]; a = official[aid]
        markers = bymarker[aid]
        bymsg = defaultdict(list)
        for h in markers:
            bymsg[h['message_id']].append(h)
        overview.append(dict(alert_id=aid, sample_role=s['sample_role'], primary_segment=s['primary_segment'],
                             diagnostic_segments=s['diagnostic_segments'], requested_segment=s.get('requested_segment'),
                             fallback=s.get('fallback', False), safety_risk_flags=s.get('safety_risk_flags', []),
                             start_local=f['start_time'].astimezone(KYIV).isoformat(), end_local=f['end_time'].astimezone(KYIV).isoformat(),
                             start_utc=f['start_time'], end_utc=f['end_time'], start_original=a['start_time_original'],
                             end_original=a['end_time_original'], duration_minutes=f['duration_minutes'],
                             v13_status=old[aid]['classification_status'], v13_categories=old[aid]['threat_combination'],
                             v14_status=f['classification_status'], v14_categories=f['threat_combination'],
                             full_timeline_message_count=len(seq[aid]), text_message_count=f['telegram_text_message_count'],
                             candidate_count=len(markers), accepted_count=sum(h['accepted_for_aggregation'] for h in markers),
                             machine_support_message_ids=f['supporting_message_ids'], evidence_strengths=f['evidence_strengths'],
                             uncertainty_marker_count=f['uncertainty_marker_count'],
                             relevant_rule_ids=sorted({rule for h in markers for rule in [h['rule_id'], *h['attribution_rule_ids']]}),
                             post_alert_diagnostic_csv='residual_post_alert_context.csv', review_status=STATUS, **blank_human()))
        windows = [('inside_interval', seq[aid], timeline),
                   ('before_alert', window(timed, times, f['start_time']-timedelta(minutes=10), f['start_time']), pre),
                   ('after_alert_diagnostic_only', window(timed, times, f['end_time'], f['end_time']+timedelta(minutes=10)), post)]
        for relation, messages, target in windows:
            reads = diagnostic_readouts(messages, dictionary, cache) if relation.startswith('after') else []
            readby = defaultdict(list)
            for h in reads:
                readby[h['message_id']].append(h)
            for n, m in enumerate(messages, 1):
                hits = bymsg[m['message_id']] if not relation.startswith('after') else []
                signals = message_signals(m, hits, dictionary)
                previous, boundaries, unresolved = contamination(aid, f['start_time'], m, data['alerts'], members) if relation == 'before_alert' else ([], [], [])
                anchors = list(dict.fromkeys(h['anchor_message_id'] for h in hits))
                cross = [dict(anchor_message_id=mid, timestamp_original=mm[mid]['message_time_original'],
                              timestamp_local=mm[mid]['message_time'].astimezone(KYIV).isoformat(), original_text=mm[mid]['text'])
                         for mid in anchors if mid != m['message_id']]
                target.append(dict(alert_id=aid, message_id=m['message_id'], chronological_index=n,
                                   time_relation=relation, timestamp_local=m['message_time'].astimezone(KYIV).isoformat(),
                                   timestamp_utc=m['message_time'], timestamp_original=m['message_time_original'],
                                   minutes_from_alert_start=(m['message_time']-f['start_time']).total_seconds()/60,
                                   original_message_text=m['text'], message_type=m['message_type'], has_text=m['has_text'],
                                   candidate_records=hits, diagnostic_post_readouts=readby[m['message_id']],
                                   rejection_reasons=[h['exclusion_reason'] for h in hits if not h['accepted_for_aggregation']],
                                   cross_message_evidence=cross, attribution_strengths=[h['attribution_strength'] for h in hits],
                                   uncertainty_reasons=[h['uncertainty_reasons'] for h in hits],
                                   diagnostic_segments=s['diagnostic_segments'], diagnostic_message_features=signals,
                                   relevant_rule_ids=sorted({rule for h in hits for rule in [h['rule_id'], *h['attribution_rule_ids']]}),
                                   previous_alert_ids=previous, intervening_boundary_alert_ids=boundaries,
                                   unresolved_envelope_alert_ids=unresolved, other_alert_memberships=sorted(members[m['message_id']]-{aid}),
                                   source_file=m['source_file'], source_message_index=m['source_message_index'],
                                   evidence_scope='available source only; not independently verified'))
    return overview, timeline, pre, post


def validate_export(selected, segments, no_text, overview, timeline, pre, post, data):
    ids = [r['alert_id'] for r in selected]
    if len(ids) != 50 or len(set(ids)) != 50:
        raise ValueError('Expected 50 unique sample alerts')
    if len(segments) != 655 or len({r['alert_id'] for r in segments}) != 655 or len(no_text) != 98:
        raise ValueError('Residual/coverage audit cardinality mismatch')
    if {r['alert_id'] for r in segments} & {r['alert_id'] for r in no_text}:
        raise ValueError('No-text cases were assigned A-F')
    expected_pairs = {(l['alert_id'], l['message_id']) for l in data['links'] if l['alert_id'] in ids}
    if len(timeline) != len(expected_pairs) or {(r['alert_id'], r['message_id']) for r in timeline} != expected_pairs:
        raise ValueError('Full timeline omitted or duplicated a message')
    mm = {m['message_id']: m for m in data['messages'] if m['message_id'] is not None}
    fm = {f['alert_id']: f for f in data['v14']}
    for s in selected:
        f = fm[s['alert_id']]
        if s['sample_role'] == 'residual' and f['classification_status'] != 'no_marker_with_text':
            raise ValueError('Residual sample is already classified')
        if s['sample_role'] != 'residual' and (f['baseline_categories'] or not f['unique_threat_category_count']):
            raise ValueError('Safety sample not newly classified')
    for rows, relation in [(timeline, 'inside'), (pre, 'pre'), (post, 'post')]:
        byalert = defaultdict(list)
        for r in rows:
            source = mm[r['message_id']]; f = fm[r['alert_id']]; t = source['message_time']
            for key, sourcekey in [('original_message_text', 'text'), ('timestamp_original', 'message_time_original'),
                                   ('timestamp_utc', 'message_time'), ('source_file', 'source_file'), ('source_message_index', 'source_message_index')]:
                if r[key] != source[sourcekey]:
                    raise ValueError('Source content/provenance changed')
            if r['timestamp_local'] != t.astimezone(KYIV).isoformat():
                raise ValueError('Incorrect Kyiv display time')
            bounds = {'inside': (f['start_time'], f['end_time']),
                      'pre': (f['start_time']-timedelta(minutes=10), f['start_time']),
                      'post': (f['end_time'], f['end_time']+timedelta(minutes=10))}
            lower, upper = bounds[relation]
            if not lower <= t < upper:
                raise ValueError('Timeline time relation invalid')
            byalert[r['alert_id']].append(source)
            for h in r['candidate_records']:
                if h['message_id'] != r['message_id'] or h['alert_id'] != r['alert_id']:
                    raise ValueError('Candidate mismatched message')
                if h['evidence_role'] == 'lexical_anchor' and source['text'][h['match_start']:h['match_end']] != h['matched_text']:
                    raise ValueError('Candidate original span invalid')
        for messages in byalert.values():
            if messages != sorted(messages, key=order):
                raise ValueError('Chronology not preserved')
        if relation != 'inside':
            timed = sorted([m for m in data['messages'] if m['message_time'] is not None and m['message_id'] is not None], key=order)
            times = [m['message_time'] for m in timed]
            for aid in ids:
                f = fm[aid]
                lower, upper = ((f['start_time']-timedelta(minutes=10), f['start_time']) if relation == 'pre'
                                else (f['end_time'], f['end_time']+timedelta(minutes=10)))
                if byalert[aid] != window(timed, times, lower, upper):
                    raise ValueError('Boundary context omitted available source messages')
    expected = [h for h in data['markers'] if h['alert_id'] in ids]
    actual = [h for r in timeline+pre for h in r['candidate_records']]
    if Counter(json.dumps(h, sort_keys=True, default=str) for h in actual) != Counter(json.dumps(h, sort_keys=True, default=str) for h in expected):
        raise ValueError('Export omitted/changed candidate records')
    for r in overview:
        if any(r[k] is not None for k in HUMAN_FIELDS):
            raise ValueError('Human decision fabricated')


def workbook(path, overview, timeline, pre, summary):
    wb = Workbook(); wb.remove(wb.active); adjustments = []
    intro = [dict(topic='Статус', explanation=STATUS),
             dict(topic='Мета', explanation='Діагностика 655 тривог із текстом. A–F — гіпотези, не підтверджені причини помилок. Окремий аудит 98 випадків без тексту у no_linked_text_source_audit.csv.'),
             dict(topic='Вибірка', explanation='40 раніше не переглянутих залишкових випадків + 10 нових класифікацій для пошуку можливого надмірного приписування. Цільовий відбір, не оцінка точності генеральної сукупності.'),
             dict(topic='Порядок', explanation='Residual Sample / False-Positive Sample → Full Timeline: прочитайте всі повідомлення по порядку, включно з no-hit. Pre-Alert Evidence — окремі доступні повідомлення за 10 хвилин до початку; не підтверджена причина тривоги.'),
             dict(topic='Граничний контекст', explanation='residual_post_alert_context.csv містить повні доступні повідомлення [end,end+10 хв). Це лише діагностична перевірка часової межі, не атрибуція та не нова категорія.'),
             dict(topic='Час', explanation='timestamp_local/start_local/end_local — ISO Europe/Kyiv з сезонним offset. UTC та оригінальні рядки збережені; H2 незмінна. DST-fold не вгадується.'),
             dict(topic='Ручне рішення', explanation='Заповнюйте Human Review: human_categories, human_decision, confirmed_failure_type, supporting_message_ids, human_comment, proposed_rule_correction. Всі поля спочатку порожні. Категорії через кому; жодне правило автоматично не впроваджується.'),
             dict(topic='Повний текст', explanation='Окремий рядок для кожного оригінального повідомлення. Якщо клітинка Excel обмежена, JSON перелічує поправки, CSV зберігає повний текст. Формули/макроси не створено.'),
             dict(topic='Сегментація', explanation='Primary priority: A,C,D,B,E,F. Страти вибірки за primary; перехресні членства збережено. Проби словника лише підсвічують фрази й не визначають тип.'),
             dict(topic='Незмінність', explanation='v1.3, v1.4, словники, джерела й Task003 залишаються незмінними. Статус технічного accepted у v1.4 не є ручним підтвердженням.')]
    readme = write_sheet(wb, 'README', intro, adjustments=adjustments)
    readme.column_dimensions['B'].width = 120
    group_rows = [dict(segment=g, overlapping_count=summary['overlapping_counts'][g],
                       primary_count=summary['primary_counts'][g], exact_condition=RULES[g]) for g in QUOTAS]
    write_sheet(wb, 'Segment Summary', group_rows, adjustments=adjustments)
    residual = [r for r in overview if r['sample_role'] == 'residual']
    safety = [r for r in overview if r['sample_role'] != 'residual']
    rs = write_sheet(wb, 'Residual Sample', residual, adjustments=adjustments)
    fs = write_sheet(wb, 'False-Positive Sample', safety, adjustments=adjustments)
    tl = write_sheet(wb, 'Full Timeline', timeline, adjustments=adjustments)
    write_sheet(wb, 'Pre-Alert Evidence', pre, adjustments=adjustments)
    hr = write_sheet(wb, 'Human Review', [dict(alert_id=r['alert_id'], sample_role=r['sample_role'], **blank_human()) for r in overview], adjustments=adjustments)
    dropdown(hr, 'human_decision', ['accept', 'reject', 'modify', 'unknown'])
    dropdown(hr, 'confirmed_failure_type', ['A', 'B', 'C', 'D', 'E', 'F', 'no_failure', 'over_attribution', 'unknown', 'other'])
    first = {}
    for n, r in enumerate(timeline, 2):
        first.setdefault(r['alert_id'], n)
    for ws in [rs, fs, hr]:
        for n in range(2, ws.max_row+1):
            aid = ws.cell(n, 1).value
            ws.cell(n, 1).hyperlink = f"#'Full Timeline'!A{first[aid]}"
            ws.cell(n, 1).font = Font(color='0563C1', underline='single')
    wb.properties.title = 'Task 004.7 — Residual uncertainty review'
    wb.properties.description = STATUS
    wb.save(path)
    loaded = load_workbook(path)
    assert loaded.sheetnames == ['README', 'Segment Summary', 'Residual Sample', 'False-Positive Sample', 'Full Timeline', 'Pre-Alert Evidence', 'Human Review']
    assert loaded['Residual Sample'].max_row == 41 and loaded['False-Positive Sample'].max_row == 11
    assert loaded['Full Timeline'].max_row == len(timeline)+1
    for ws in [loaded['Human Review'], loaded['Residual Sample'], loaded['False-Positive Sample']]:
        heads = [c.value for c in ws[1]]
        for field in HUMAN_FIELDS:
            index = heads.index(field)+1
            assert all(ws.cell(n, index).value is None for n in range(2, ws.max_row+1))
    assert not any(c.data_type == 'f' for ws in loaded for row in ws for c in row)
    loaded.close()
    return adjustments


def markdown(report):
    lines = ['# Task 004.7 — Residual uncertainty segmentation', '', '**'+STATUS+'**', '',
             'Frozen baseline: '+json.dumps(report['baseline'])+'.', '',
             'A–F are diagnostic hypotheses, not validated root causes. Membership overlaps; primary groups partition all 655 text residuals.', '',
             '| Segment | Overlapping | Primary |', '|---|---:|---:|']
    for g in QUOTAS:
        lines.append(f"| {g} | {report['overlapping_counts'][g]} | {report['primary_counts'][g]} |")
    lines += ['', '## Exact diagnostic rules', 'Primary priority: '+','.join(PRIMARY_PRIORITY)+'.']
    lines += ['- '+g+': '+RULES[g] for g in QUOTAS]
    lines += ['', 'Diagnostic lexical probes (not classifier rules): '+json.dumps(GAP_PROBES, ensure_ascii=False), '',
              '## Source coverage audit', '98 no-linked-text cases stay outside A–F. Subtypes: '+json.dumps(report['source_gap_counts'])+'. The audit retains nearest before/after message IDs and posting gaps; no threat absence is inferred.', '',
              '## Controlled selection', 'Seed 20261011. Exclude all 25 Task0045 sample alerts from both new samples and the five discussed cases. Within each primary group: lexicographically sorted IDs, seeded per-year shuffle, cycle available years in ascending order; within each next year greedily prefer unseen recorded rejection/probe/geography/boundary/volume signals, seeded order resolves ties. Requested A8/B8/C8/D6/E5/F5. If short: report shortfall, fill from unused eligible cases with the same year-balanced method, retain actual primary and fallback flag.',
              'Quota audit: '+json.dumps(report['selection_audit']),
              'Safety sample: newly classified only; seeded tie order, greedy unseen priority-risk flags, then fewer selected in year, then number of flags. Risks are not proof of false positives. All selectors are purposive; no accuracy estimation.', '']
    for g in QUOTAS:
        lines.append(g+': '+', '.join(report['selected_residual_by_primary'][g]))
    lines += ['Safety: '+', '.join(report['selected_safety_ids']), '', '## Evidence completeness',
              f"50 unique alerts; {report['timeline_messages']} complete linked in-alert messages, including {report['no_hit_timeline_messages']} with no lexical anchors; {report['pre_alert_messages']} pre-alert messages; {report['post_alert_messages']} post-alert diagnostic-only messages. Candidate occurrences: {report['candidate_records']}. All human fields blank.",
              'All chronology rows preserve original text/title/ID/file/index. Half-open links are Task003 links, not recomputed assignments. Pre-alert and post-alert windows each retrieve every available timed record for ten minutes. Post-alert lexical readouts reuse the frozen profile, are explicitly unaccepted, and are only diagnostic boundary checks. Cross-message records cite original anchors and their unmodified texts. Never infer continuity across unrelated messages.', '',
              '## Limitations', 'Recorded-signal E is not a semantic majority judgment. Place-only B cannot establish weapon type. C probes may be benign, unrelated, or already understood semantically; absence of a hit is not a dictionary error. D proximity can concern a different episode or deliberate safe rejection. Primary priority affects counts and sampling. Missing media/OCR, one-channel coverage, untimed/DST-ambiguous records and unrecorded legacy context/geography regex IDs limit explanation. v1.4 uncertainty remains explicit. October official starts stop Oct4; Telegram extends Oct10. No classifier or approved interpretation is changed.', '',
              '## Validation', 'Validation: '+json.dumps(report['validation'])+'. Protected data/config/classifier/prior-report hashes verified unchanged. Tests and reproducibility checks are recorded separately in validation_results.json.',
              f"Excel cell representation adjustments: {len(report['excel_cell_adjustments'])}; full details in JSON, exact texts in CSV.", '',
              '## Reproduction', '`python -m kyiv_alerts.residual_review --root .`', '',
              'Outputs: residual_segment_summary.md/.json; residual_alert_segments.csv; residual_review_50.xlsx; residual_review_alerts.csv; residual_full_timeline.csv; residual_pre_alert_evidence.csv; false_positive_review_sample.csv. Supplements: no_linked_text_source_audit.csv, residual_post_alert_context.csv, residual_human_review.csv.', '',
              'Stop for dialogue-based human validation. v1.5, time-series analysis and forecasting are not implemented.']
    return '\n'.join(lines)+'\n'


def run(root):
    tick = time.perf_counter(); root = root.resolve()
    protected = snapshot(root)
    data = {key: pq.read_table(root/path).to_pylist() for key, path in FILES.items()}
    counts = baseline(data['v14'])  # Fail before creating outputs if baseline differs.
    validate_h2(data['messages'])
    validate_links(data['links'], data['coverage'], data['alerts'], data['messages'])
    if len({r['alert_id'] for r in data['v14']}) != counts['official']:
        raise ValueError('Official feature IDs not unique')
    official = {a['alert_id']: a for a in data['alerts']}
    for version in ['v13', 'v14']:
        if len(data[version]) != len(official) or {f['alert_id'] for f in data[version]} != set(official):
            raise ValueError('Official IDs differ in '+version)
        for f in data[version]:
            if any(f[k] != official[f['alert_id']][k] for k in ['start_time', 'end_time', 'duration_minutes']):
                raise ValueError('Official interval changed in '+version)
    previous_report = json.loads((root/'reports/classification_v1_4/classification_comparison.json').read_text())
    for rel, expected in previous_report['protected_hashes'].items():
        if digest(root/rel) != expected:
            raise ValueError('Frozen historical source changed: '+rel)
    policy = yaml.safe_load((root/'config/attribution_v1_4.yaml').read_text())
    dictionary = LexicalProfile(root/'config/threat_dictionary_v1_3.yaml', root/'config/kyiv_geography_v1_1.yaml', policy)
    mm = {m['message_id']: m for m in data['messages'] if m['message_id'] is not None}
    seq, bymarker = defaultdict(list), defaultdict(list)
    for l in data['links']:
        seq[l['alert_id']].append(mm[l['message_id']])
    for aid in seq:
        seq[aid].sort(key=order)
    for h in data['markers']:
        bymarker[h['alert_id']].append(h)
    timed = sorted([m for m in data['messages'] if m['message_time'] is not None and m['message_id'] is not None], key=order)
    times = [m['message_time'] for m in timed]; cache = {}
    segments = []
    for f in data['v14']:
        if f['classification_status'] != 'no_marker_with_text':
            continue
        posts = window(timed, times, f['end_time'], f['end_time']+timedelta(minutes=10))
        segments.append(segment(f, seq[f['alert_id']], bymarker[f['alert_id']], dictionary,
                                diagnostic_readouts(posts, dictionary, cache)))
    with (root/'reports/controlled_review/selected_alert_overview.csv').open(encoding='utf-8-sig', newline='') as stream:
        reviewed = {r['alert_id'] for r in csv.DictReader(stream)}
    if len(reviewed) != 25:
        raise ValueError('Task0045 sample is not the expected 25')
    discussed = {'2563', '2565', '2581', '2578', '2590'}
    residual, audit = select_residual(segments, reviewed | discussed)
    safety = select_safety(data['v14'], bymarker, reviewed | discussed)
    selected = residual+safety
    no_text = source_coverage_audit(data['v14'], seq, timed, times)
    overview, timeline, pre, post = reconstruct(selected, data, seq, bymarker, timed, times, dictionary, cache)
    validate_export(selected, segments, no_text, overview, timeline, pre, post, data)
    overlaps = {g: sum(g in s['diagnostic_segments'] for s in segments) for g in QUOTAS}
    primaries = {g: sum(s['primary_segment'] == g for s in segments) for g in QUOTAS}
    assert sum(primaries.values()) == 655 and len(residual) == 40 and len(safety) == 10
    assert not {r['alert_id'] for r in selected} & (reviewed | discussed)
    report = dict(status=STATUS, baseline=counts, seed=SEED, overlapping_counts=overlaps, primary_counts=primaries,
                  primary_priority=PRIMARY_PRIORITY, grouping_rules=RULES, diagnostic_gap_probes=GAP_PROBES,
                  selection_audit=audit, excluded_task0045_ids=sorted(reviewed), excluded_discussed_ids=sorted(discussed),
                  selected_residual_by_primary={g: [r['alert_id'] for r in residual if r['primary_segment'] == g] for g in QUOTAS},
                  selected_safety_ids=[r['alert_id'] for r in safety],
                  source_gap_counts=dict(Counter(r['source_gap_type'] for r in no_text)),
                  overlapping_combinations=dict(Counter('+'.join(r['diagnostic_segments']) for r in segments)),
                  year_primary_counts={str(y): {g: sum(r['year'] == y and r['primary_segment'] == g for r in segments) for g in QUOTAS} for y in range(2023, 2027)},
                  timeline_messages=len(timeline), no_hit_timeline_messages=sum(not any(h['evidence_role'] == 'lexical_anchor' for h in r['candidate_records']) for r in timeline),
                  pre_alert_messages=len(pre), post_alert_messages=len(post),
                  candidate_records=sum(len(r['candidate_records']) for r in timeline+pre),
                  protected_hashes=protected, input_hashes={k: digest(root/p) for k, p in FILES.items()},
                  validation=dict(complete_source_chronology=True, correct_links_and_local_times=True, all_human_fields_blank=True,
                                  source_and_classifier_hashes_unchanged=True, unique_samples=True, no_text_separate=True),
                  excel_cell_adjustments=[])
    out = root/OUTPUT; out.mkdir(parents=True, exist_ok=True)
    for name, rows in [('residual_alert_segments', segments), ('residual_review_alerts', overview),
                       ('residual_full_timeline', timeline), ('residual_pre_alert_evidence', pre),
                       ('false_positive_review_sample', [r for r in overview if r['sample_role'] != 'residual']),
                       ('no_linked_text_source_audit', no_text), ('residual_post_alert_context', post),
                       ('residual_human_review', [dict(alert_id=r['alert_id'], sample_role=r['sample_role'], **blank_human()) for r in overview])]:
        write_audit_csv(rows, out/(name+'.csv'))
    report['excel_cell_adjustments'] = workbook(out/'residual_review_50.xlsx', overview, timeline, pre, report)
    if snapshot(root) != protected:
        raise ValueError('Protected source/classifier/prior report was mutated')
    report['elapsed_seconds'] = round(time.perf_counter()-tick, 6)
    (out/'residual_segment_summary.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    (out/'residual_segment_summary.md').write_text(markdown(report), encoding='utf-8')
    print(json.dumps({k: report[k] for k in ['baseline', 'overlapping_counts', 'primary_counts', 'selection_audit', 'selected_residual_by_primary', 'selected_safety_ids', 'source_gap_counts', 'timeline_messages', 'pre_alert_messages', 'post_alert_messages', 'elapsed_seconds']}, ensure_ascii=False))
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', type=Path, default=Path('.'))
    run(parser.parse_args().root)
