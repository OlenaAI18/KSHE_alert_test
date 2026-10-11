"""Full-data isolated v1.5 comparison, regression reconciliation and audit export."""
import argparse
from collections import Counter, defaultdict
import copy
import csv
import json
from pathlib import Path
import time

import pyarrow as pa
import pyarrow.parquet as pq
import yaml
from openpyxl import Workbook, load_workbook

from .core import process, key, STATUS
from .lexical import Profile
from kyiv_alerts.residual_review.__main__ import baseline, snapshot as old_snapshot, window
from kyiv_alerts.residual_review.core import balanced_pick, blank_human
from kyiv_alerts.controlled_review.__main__ import digest
from kyiv_alerts.controlled_review.core import order
from kyiv_alerts.controlled_review.workbook import write_sheet, dropdown
from kyiv_alerts.classification.consolidation import summary, compare_versions, temporal_summaries, CATEGORIES
from kyiv_alerts.classification.coverage_audit import write_audit_csv
from kyiv_alerts.integration.__main__ import validate_h2, validate_links
from kyiv_alerts.ingestion.timeutils import KYIV
from datetime import timedelta
import random

DATA = Path('data/processed/classification_v1_5')
REPORT = Path('reports/classification_v1_5')
FILES = {'alerts':'data/interim/alerts.parquet', 'messages':'data/interim/telegram_messages.parquet',
         'links':'data/processed/alert_message_links.parquet', 'coverage':'data/processed/alert_coverage.parquet',
         'v13':'data/processed/classification_v1_3/alert_threat_features.parquet',
         'v14':'data/processed/classification_v1_4/alert_threat_features.parquet',
         'markers':'data/processed/classification_v1_4/message_threat_markers.parquet'}
CAUSES = ('A','B','C','D','scope')


def snapshot(root):
    result = {k:v for k,v in old_snapshot(root).items()
            if not k.startswith((str(DATA)+'/', str(REPORT)+'/'))
            and k not in {'config/attribution_v1_5.yaml', 'config/task0048_reviewed_cases.yaml'}}
    for folder in ['reports/classification_v1_4/residual_review', 'src/kyiv_alerts']:
        for p in sorted((root/folder).rglob('*')):
            if p.is_file() and '__pycache__' not in p.parts and not p.is_relative_to(root/'src/kyiv_alerts/classification_v15'):
                result[str(p.relative_to(root))] = digest(p)
    return result


def rule_hashes(root):
    return {str(p.relative_to(root)):digest(p) for p in sorted((root/'src/kyiv_alerts/classification_v15').glob('*.py'))} | {
        'config/attribution_v1_5.yaml':digest(root/'config/attribution_v1_5.yaml')}


def evidence_identity(row):
    return key(row)+(row['category'],row['anchor_message_id'],row['anchor_rule_id'],
                     row['anchor_match_start'],row['anchor_match_end'])


def validate(markers, features, data):
    fm = {f['alert_id']:f for f in features}; old = {f['alert_id']:f for f in data['v14']}
    mm = {m['message_id']:m for m in data['messages'] if m['message_id'] is not None}
    pairs = {(r['alert_id'],r['message_id']) for r in data['links']}; accepted = defaultdict(list)
    if len(fm)!=2594 or len(features)!=2594 or set(fm)!=set(old):
        raise ValueError('Official alert cardinality/identity changed')
    identities = set()
    anchors = {(m['alert_id'],m['message_id'],m['rule_id'],m['match_start'],m['match_end'],m['category']):m
               for m in markers if m['evidence_role']=='lexical_anchor'}
    for m in markers:
        identity = evidence_identity(m)
        if identity in identities: raise ValueError('Duplicate evidence identity')
        identities.add(identity)
        source=mm[m['message_id']];f=fm[m['alert_id']]
        if m['category'] not in CATEGORIES[:-1]: raise ValueError('Threat taxonomy expanded')
        for k,v in [('original_text',source['text']),('source_timestamp_original',source['message_time_original']),
                    ('source_message_time',source['message_time']),('source_file',source['source_file']),('source_message_index',source['source_message_index'])]:
            if m[k]!=v: raise ValueError('Evidence source provenance changed: '+k)
        for trace in m['scope_trace']:
            if source['text'][trace['match_start']:trace['match_end']]!=trace['matched_expression']:
                raise ValueError('Scoped exclusion span not original source')
        if m['evidence_role']=='lexical_anchor' and source['text'][m['match_start']:m['match_end']]!=m['matched_text']:
            raise ValueError('Lexical source span changed')
        if m['time_relation']=='inside_interval' and (m['alert_id'],m['message_id']) not in pairs:
            raise ValueError('Inside evidence not frozen Task003 link')
        if m['time_relation']=='before_alert':
            delta=(f['start_time']-source['message_time']).total_seconds()/60
            if not 0<delta<=10 or m['minutes_before_alert']!=delta: raise ValueError('Pre-alert boundary invalid')
        if m['evidence_role']=='context_reference':
            anchor=anchors[(m['alert_id'],m['anchor_message_id'],m['anchor_rule_id'],m['anchor_match_start'],m['anchor_match_end'],m['category'])]
            if m['accepted_for_aggregation'] and not anchor['accepted_for_aggregation']:
                raise ValueError('Accepted reference has no accepted explicit anchor')
            if not 0<=(source['message_time']-mm[m['anchor_message_id']]['message_time']).total_seconds()<=600:
                raise ValueError('Episode exceeded ten-minute anchor age')
        if m['accepted_for_aggregation']:
            if m['previous_alert_ids'] or m['intervening_boundary_alert_ids'] or m['unresolved_envelope_alert_ids']:
                raise ValueError('Accepted contaminated evidence')
            if any(t['blocks_occurrence'] for t in m['scope_trace']): raise ValueError('Scoped negation accepted')
            if not m['v14_accepted'] and not m['uncertainty_flag']: raise ValueError('New evidence uncertainty lost')
            if m['time_relation']=='before_alert' and not any((m['alert_id'],mid) in pairs for mid in m['support_message_ids']):
                raise ValueError('Pre-alert warning accepted alone')
            accepted[m['alert_id']].append(m)
    for aid,f in fm.items():
        if any(f[k]!=old[aid][k] for k in ['start_time','end_time','duration_minutes','coverage_status','historical_range_relation','telegram_message_count','telegram_text_message_count']):
            raise ValueError('Original temporal/coverage metadata changed')
        cats={m['category'] for m in accepted[aid]}
        if cats!={c for c in CATEGORIES[:-1] if f[c] is True}: raise ValueError('Feature/marker category mismatch')
        if old[aid]['classification_status'] in {'outside_telegram_coverage','unresolved'}:
            if cats or any(f[c] is not None for c in CATEGORIES[:-1]): raise ValueError('Unavailable observation fabricated')
        ids=f['supporting_message_ids']
        if ids is not None:
            expected=sorted({mid for m in accepted[aid] for mid in m['support_message_ids']},key=lambda mid:order(mm[mid]))
            if ids!=expected: raise ValueError('Source support IDs not chronological/deduplicated')


def regressions(root, features, markers, data):
    old_cases=yaml.safe_load((root/'config/task0046_reviewed_cases.yaml').read_text())['review_cases']
    new_cases=yaml.safe_load((root/'config/task0048_reviewed_cases.yaml').read_text())['review_cases']
    cases={r['alert_id']:r for r in old_cases+new_cases}
    with (root/'reports/controlled_review/selected_alert_overview.csv').open(encoding='utf-8-sig',newline='') as stream:
        prior=list(csv.DictReader(stream))
    priorids={r['alert_id'] for r in prior}; ids=sorted(priorids|set(cases))
    fm={f['alert_id']:f for f in features};v14={f['alert_id']:f for f in data['v14']};by=defaultdict(list)
    for m in markers:
        if m['accepted_for_aggregation']:by[m['alert_id']].append(m)
    result=[]
    for aid in ids:
        case=cases.get(aid); observed=[c for c in CATEGORIES[:-1] if fm[aid][c] is True]
        expected=case['approved_categories'] if case else None
        missing=sorted(set(expected or [])-set(observed));extra=sorted(set(observed)-set(expected or [])) if case else []
        status='NO_CONFIRMED_HUMAN_LABEL' if case is None else 'PROVISIONAL_EXAMPLE' if 'provisional' in case['label_status'] else 'CONFIRMED_LABEL_DISCREPANCY' if missing or extra else 'CONFIRMED_LABEL_REPRODUCED'
        result.append(dict(alert_id=aid, task0045_sample=aid in priorids,
                           label_status=case['label_status'] if case else 'not_human_labeled',
                           expected_categories=expected, v14_categories=v14[aid]['threat_combination'],
                           v15_categories=observed, missing_expected_categories=missing,
                           additional_experimental_categories=extra, regression_status=status,
                           existing_research_lead_decision=case.get('research_lead_decision') if case else None,
                           supporting_evidence=[dict(message_id=m['message_id'],anchor_message_id=m['anchor_message_id'],
                               original_text=m['original_text'],timestamp_original=m['source_timestamp_original'],
                               time_relation=m['time_relation'],category=m['category'],strength=m['attribution_strength'],
                               rules=m['attribution_rule_ids'],uncertainty=m['uncertainty_reasons']) for m in by[aid]],
                           note='User reconfirmed two-category target; pre-alert ballistic warning lacks type-linked in-alert corroboration. No ID exception introduced.' if aid=='1096' and 'ballistic_high_speed' in missing else case.get('clarification') if case else 'Task0045 review fields remain blank; no gold label inferred.'))
    return result


def audit_sample(features, new_ids, excluded, bymarker):
    pools=[]; selected=[]; used=set();audit=[];rng=random.Random(20261011)
    eligible=[dict(alert_id=f['alert_id'],year=f['start_time'].astimezone(KYIV).year,
                   sampling_signal_types=(f['correction_causes'] or [])+['category.'+c for c in CATEGORIES[:-1] if f[c] is True])
              for f in features if f['alert_id'] in new_ids and f['alert_id'] not in excluded]
    fm={f['alert_id']:f for f in features}
    conditions=[('strategic_aviation',lambda f:f['strategic_aviation'] is True),
                ('ballistic_high_speed',lambda f:f['ballistic_high_speed'] is True),
                ('pre_alert_or_cross_message',lambda f: f['has_pre_alert_evidence'] or any(c in {'C','D'} for c in f['correction_causes'] or [])),
                ('other_uncertain',lambda f:bool(f['uncertainty_marker_count']))]
    for name, condition in conditions:
        pool=[r for r in eligible if r['alert_id'] not in used and condition(fm[r['alert_id']])]
        chosen=balanced_pick(pool,5,rng)
        selected += [dict(r,audit_stratum=name,fallback=False) for r in chosen]
        used.update(r['alert_id'] for r in chosen)
        audit.append(dict(stratum=name,requested=5,available=len(pool),selected=len(chosen),shortfall=5-len(chosen)))
    selected += [dict(r,audit_stratum='fallback_weak_evidence',fallback=True) for r in balanced_pick([r for r in eligible if r['alert_id'] not in used],20-len(selected),rng)]
    return selected,audit


def export_audit(root, features, markers, selected, data):
    fm={f['alert_id']:f for f in features};mm={m['message_id']:m for m in data['messages'] if m['message_id'] is not None}
    seq=defaultdict(list);bm=defaultdict(list)
    for l in data['links']:seq[l['alert_id']].append(mm[l['message_id']])
    for m in markers:bm[(m['alert_id'],m['message_id'])].append(m)
    timed=sorted([m for m in data['messages'] if m['message_time'] is not None and m['message_id'] is not None],key=order)
    times=[m['message_time'] for m in timed];overview=[];timeline=[];pre=[]
    for s in selected:
        aid=s['alert_id'];f=fm[aid]
        overview.append(dict(alert_id=aid,audit_stratum=s['audit_stratum'],fallback=s['fallback'],
                             start_local=f['start_time'].astimezone(KYIV).isoformat(),end_local=f['end_time'].astimezone(KYIV).isoformat(),
                             duration_minutes=f['duration_minutes'],v14_categories=f['baseline_categories'],
                             v15_categories=f['threat_combination'],evidence_strengths=f['evidence_strengths'],
                             correction_causes=f['correction_causes'],uncertainty_count=f['uncertainty_marker_count'],
                             review_status=STATUS,**blank_human()))
        for relation, messages, output in [('inside_interval',sorted(seq[aid],key=order),timeline),
                ('before_alert',window(timed,times,f['start_time']-timedelta(minutes=10),f['start_time']),pre)]:
            for n,m in enumerate(messages,1):
                hits=bm[(aid,m['message_id'])]
                output.append(dict(alert_id=aid,message_id=m['message_id'],chronological_index=n,time_relation=relation,
                                   timestamp_local=m['message_time'].astimezone(KYIV).isoformat(),timestamp_utc=m['message_time'],
                                   timestamp_original=m['message_time_original'],original_message_text=m['text'],
                                   candidate_records=hits,reasons_for_attribution=[h['attribution_rule_ids'] for h in hits if h['accepted_for_aggregation']],
                                   cross_message_anchors=[dict(message_id=h['anchor_message_id'],original_text=mm[h['anchor_message_id']]['text'],
                                                               timestamp_original=mm[h['anchor_message_id']]['message_time_original']) for h in hits if h['anchor_message_id']!=m['message_id']],
                                   source_file=m['source_file'],source_message_index=m['source_message_index']))
    ids={s['alert_id'] for s in selected}
    assert len(ids)==20 and len(selected)==20
    assert {(r['alert_id'],r['message_id']) for r in timeline}=={(l['alert_id'],l['message_id']) for l in data['links'] if l['alert_id'] in ids}
    wb=Workbook();wb.remove(wb.active);adjust=[]
    write_sheet(wb,'README',[dict(topic='Status',explanation=STATUS),
        dict(topic='Purpose',explanation='20 newly classified alerts, separate from development and reserved holdout. Purposive stratified weak-evidence audit, not population accuracy estimation.'),
        dict(topic='Instructions',explanation='Read every Full Timeline row and separate Pre-Alert Evidence. Human Review fields start blank. No-hit/empty records remain visible. Local ISO times include Kyiv seasonal offsets; original titles remain unmodified.'),
        dict(topic='Limits',explanation='CSV retains full original text if Excel limits require a flagged representation adjustment. No categories are confirmed by temporal association alone.')],adjustments=adjust)
    write_sheet(wb,'Audit Sample',overview,adjustments=adjust)
    write_sheet(wb,'Full Timeline',timeline,adjustments=adjust)
    write_sheet(wb,'Pre-Alert Evidence',pre,adjustments=adjust)
    hr=write_sheet(wb,'Human Review',[dict(alert_id=r['alert_id'],**blank_human()) for r in overview],adjustments=adjust)
    dropdown(hr,'human_decision',['accept','reject','modify','unknown'])
    out=root/REPORT
    for name,rows in [('false_positive_audit_20',overview),('audit_full_timeline',timeline),('audit_pre_alert_evidence',pre)]:
        write_audit_csv(rows,out/(name+'.csv'))
    wb.save(out/'false_positive_audit_20.xlsx')
    check=load_workbook(out/'false_positive_audit_20.xlsx')
    assert check['Audit Sample'].max_row==21 and check['Human Review'].max_row==21
    assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
    assert all(check['Human Review'].cell(r,c).value is None for r in range(2,22) for c in range(2,8))
    check.close()
    return dict(selected_ids=[r['alert_id'] for r in selected],timeline_messages=len(timeline),pre_alert_messages=len(pre),excel_cell_adjustments=adjust)


def markdown(r):
    lines=['# Task 004.8 — Targeted attribution v1.5','', '**'+STATUS+'**',
           'Technical status: '+r['task_status']+'. Human-label discrepancies are not silently overridden.', '',
           '| Version | Classified | Eligible coverage | Text residual |', '|---|---:|---:|---:|']
    for v in ['v1_3','v1_4','v1_5']:
        s=r[v];lines.append(f"| {v} | {s['classified_alerts']} | {s['any_marker_yield_eligible']:.2%} | {s['no_marker_with_text']} |")
    c=r['comparison_v14_v15']
    lines += ['',f"Versus v1.4: gained {c['gained_count']} unique alerts; lost {c['lost_count']}; net {c['gained_count']-c['lost_count']}. Original 2594 rows / 1803 eligible remain. No-text/outside/DST states remain separately audited.", '',
              '| Category | v1.3 | v1.4 | v1.5 | Delta vs v1.4 |', '|---|---:|---:|---:|---:|']
    for cat in CATEGORIES[:-1]:
        counts=[r[v]['per_category_alert_counts'][cat] for v in ['v1_3','v1_4','v1_5']]
        lines.append(f"| {cat} | {counts[0]} | {counts[1]} | {counts[2]} | {counts[2]-counts[1]} |")
    lines += ['', '## Corrections and scope',
        'A: full operational statements with occurrence-specific context; introductory question clauses do not negate following statements. B: narrowly composed ballistic-weaponry and enemy-UAV expressions; no taxonomy expansion. C: bounded current episodes may reconsider a route toward Kyiv region; static unrelated regional threats cannot jump regions. D: pre-alert lexical anchors require in-alert correspondence, not a warning alone. Ordinary gap remains 5min and anchor age 10min; an explicit Kyiv alert-cause pronoun may bridge a >5min pre/inside gap only within the same 10min anchor-age bound. Multiple simultaneous families leave a pronoun unresolved.',
        'Scope: explicit negation/history/quotation/explanation retains exact original source spans. Movement reversal is not cancellation. Lost observation stops propagation but does not delete previously observed evidence. No strikes is not no UAV. Negation of one family does not cancel another. Valid v1.4 evidence is retained; newly invalid scoped evidence and unsupported pre-only transfers are withdrawn with traces.', '',
        'Overlapping rule supports: '+json.dumps(r['rule_gain_breakdown']),
        'Full-data leave-one-correction-out ablations: '+json.dumps(r['rule_ablations']),
        'Support and ablation counts overlap and must not be added. Ablations always retain safety guards; scope ablation disables rescuing legacy false cancellations, not negation safety.', '',
        '## Reviewed cases']
    for case in r['regression_cases']:
        if case['expected_categories'] is not None:
            lines.append(f"- {case['alert_id']}: target {case['expected_categories']}; observed {case['v15_categories']}; {case['label_status']}; {case['regression_status']}. {case['note'] or ''}")
    lines += ['All 25 Task0045 sample cases are included in regression_cases.csv. Blank prior human review fields are not inferred gold labels. 1416 is provisional; 2590 remains a working hypothesis. The prior 2581 discrepancy remains explicitly retained. For 1096 the user reconfirmed shahed_uav + ballistic_high_speed, but the ballistic pre-warning is not type-linked to inside evidence under the required policy. It remains a confirmed-label/source-attribution discrepancy, not a category forced by ID.', '',
        '## Month/year and recent coverage', '| Month | v1.3 | v1.4 | v1.5 | Eligible |', '|---|---:|---:|---:|---:|']
    for x in r['recent_months']:
        lines.append(f"| {x['month']} | {x['v13_classified']} | {x['v14_classified']} | {x['v15_classified']} | {x['eligible']} |")
    lines += ['All month/year rows are CSV/JSON; October official starts end Oct4, last interval ends Oct5, Telegram extends Oct10.', '',
        '## Evidence and audit', 'Unique-alert strength distribution (overlapping): '+json.dumps(r['evidence_strength_alerts']),
        'Pre-alert-only lexical categories with in-alert corroboration: '+json.dumps(r['pre_alert_only_lexical']),
        'Newly classified uncertain alerts: '+str(r['new_uncertain_classifications'])+'. Unsupported pre-warning-alone accepted alerts: 0.',
        'Reserved diagnostic holdout: '+str(len(r['holdout']['reserved_alert_ids']))+' cases. No textual tuning or inferred human labels. Per-ID category outputs are generated only after frozen rule hashes; holdout remains diagnostic and unvalidated.',
        'Separate twenty-case audit: '+json.dumps(r['audit']), '',
        '## Risks and validation',
        'Broad UAV category is not a confirmed Shahed model. A ballistic launch origin and a nearby route are not independent proof of Kyiv direction. Short-pronoun identity, coincident objects, geography ambiguity, unrecorded legacy regex trace and one-channel coverage remain uncertain. New rule gains are coverage, not accuracy. Review the 20 cases and confirmed discrepancies before adoption.',
        'Source/data/config/frozen-classifier/prior-report hashes verified unchanged; full source FK/span/time/half-open membership and feature/marker reconciliation checks passed. Runtime and checks are recorded in JSON and validation_results.json. No forecasting, baseline overwrite or automatic commit.', '',
        'Reproduce: `python -m kyiv_alerts.classification_v15 --root .`. Stop after Task0048.']
    return '\n'.join(lines)+'\n'


def run(root):
    root=root.resolve();tick=time.perf_counter();protected=snapshot(root);frozen_rules=rule_hashes(root)
    data={k:pq.read_table(root/p).to_pylist() for k,p in FILES.items()}
    try:
        baseline(data['v14'])
        if abs(summary(data['v14'])['any_marker_yield_eligible']*100-58.24)>0.005:
            raise ValueError('Baseline coverage differs')
    except ValueError as error:
        out=root/REPORT;out.mkdir(parents=True,exist_ok=True)
        (out/'baseline_abort.json').write_text(json.dumps(dict(status='BLOCKED',reason=str(error)),indent=2)+'\n')
        raise
    validate_h2(data['messages']);validate_links(data['links'],data['coverage'],data['alerts'],data['messages'])
    historical=json.loads((root/'reports/classification_v1_4/residual_review/residual_segment_summary.json').read_text())
    for path,value in historical['protected_hashes'].items():
        if digest(root/path)!=value:raise ValueError('Frozen historical artifact changed: '+path)
    policy=yaml.safe_load((root/'config/attribution_v1_5.yaml').read_text())
    policy14=yaml.safe_load((root/'config/attribution_v1_4.yaml').read_text())
    dictionary=Profile(root/'config/threat_dictionary_v1_3.yaml',root/'config/kyiv_geography_v1_1.yaml',policy14,policy)
    development=yaml.safe_load((root/'config/task0048_reviewed_cases.yaml').read_text())['review_cases']
    development_ids={r['alert_id'] for r in development}
    all_review={aid for rows in historical['selected_residual_by_primary'].values() for aid in rows}|set(historical['selected_safety_ids'])
    holdout=dict(status='UNTOUCHED_DIAGNOSTIC_HOLDOUT',reserved_alert_ids=sorted(all_review-development_ids),
                 development_ids=sorted(development_ids),no_textual_tuning=True,frozen_rule_hashes=frozen_rules)
    cache={};markers,features=process(data['v14'],data['messages'],data['links'],data['markers'],dictionary,data['alerts'],policy,cache)
    validate(markers,features,data)
    comparison,changes=compare_versions(data['v14'],features)
    comp13,_=compare_versions(data['v13'],features)
    gained=set(comparison['gained_alert_ids']);lost=set(comparison['lost_alert_ids'])
    by=defaultdict(list)
    for m in markers:by[m['alert_id']].append(m)
    supports=[]
    for cause in CAUSES:
        aids={m['alert_id'] for m in markers if m['accepted_for_aggregation'] and cause in m['correction_causes']}
        supports.append(dict(correction=cause,newly_classified_supported_unique_alerts=len(aids&gained),
                             all_supported_unique_alerts=len(aids),supported_gained_ids=sorted(aids&gained)))
    ablations=[]
    for cause in CAUSES:
        alternate=copy.deepcopy(policy);alternate['enabled_corrections']=[c for c in CAUSES if c!=cause]
        engine=dictionary if cause!='B' else Profile(root/'config/threat_dictionary_v1_3.yaml',root/'config/kyiv_geography_v1_1.yaml',policy14,alternate)
        _,ff=process(data['v14'],data['messages'],data['links'],data['markers'],engine,data['alerts'],alternate,cache if cause!='B' else {})
        ids={f['alert_id'] for f in ff if not f['unknown']};full={f['alert_id'] for f in features if not f['unknown']}
        ablations.append(dict(disabled_correction=cause,classified_alerts=len(ids),
                              unique_alerts_lost_when_disabled=len(full-ids),unique_alerts_gained_when_disabled=len(ids-full)))
    regression=regressions(root,features,markers,data)
    excluded=set(holdout['reserved_alert_ids'])|{r['alert_id'] for r in regression}
    selected,audit_strata=audit_sample(features,gained,excluded,by)
    if len(selected)!=20: raise ValueError('Insufficient newly classified non-holdout cases for requested 20 audit')
    month13,year13=temporal_summaries(data['v13']);month14,year14=temporal_summaries(data['v14']);month15,year15=temporal_summaries(features)
    def calendars(a,b,c,label):
        return [dict(period=x[label],**{label:x[label]},eligible=x['eligible_alerts'],
                     v13_classified=u['classified_alerts'],v14_classified=v['classified_alerts'],v15_classified=x['classified_alerts'],
                     v13_coverage=u['any_marker_yield_eligible'],v14_coverage=v['any_marker_yield_eligible'],v15_coverage=x['any_marker_yield_eligible']) for u,v,x in zip(a,b,c)]
    months=calendars(month13,month14,month15,'month');years=calendars(year13,year14,year15,'year')
    current={evidence_identity(m):m for m in markers};withdrawn=[]
    for old in data['markers']:
        if not old['accepted_for_aggregation']:continue
        now=current.get(evidence_identity(old))
        if now is None or not now['accepted_for_aggregation']:
            withdrawn.append(dict(alert_id=old['alert_id'],message_id=old['message_id'],rule_id=old['rule_id'],category=old['category'],
                                  evidence_role=old['evidence_role'],anchor_message_id=old['anchor_message_id'],
                                  anchor_rule_id=old['anchor_rule_id'],anchor_match_start=old['anchor_match_start'],anchor_match_end=old['anchor_match_end'],
                                  matched_text=old['matched_text'],previous_strength=old['attribution_strength'],
                                  current_exclusion=now['exclusion_reason'] if now else 'context_reference_not_reproduced',
                                  scope_trace=now['scope_trace'] if now else [],original_text=old['original_text'],
                                  lost_alert=old['alert_id'] in lost))
    pre_only=[]
    for f in features:
        for c in CATEGORIES[:-1]:
            rows=[m for m in by[f['alert_id']] if m['accepted_for_aggregation'] and m['category']==c]
            anchors=[m for m in rows if m['evidence_role']=='lexical_anchor']
            if anchors and all(m['time_relation']=='before_alert' for m in anchors):
                pre_only.append(dict(alert_id=f['alert_id'],category=c,support_message_ids=f['supporting_message_ids'],
                                     inside_corroboration_required=True))
    source_ids={m['message_id']:m for m in data['messages'] if m['message_id'] is not None}
    out=root/REPORT;out.mkdir(parents=True,exist_ok=True);folder=root/DATA;folder.mkdir(parents=True,exist_ok=True)
    for name,rows in [('message_threat_markers',markers),('alert_threat_features',features)]:
        table=pa.Table.from_pylist(rows);pq.write_table(table,folder/(name+'.parquet'),compression='zstd')
        assert pq.read_table(folder/(name+'.parquet')).equals(table)
    for name,rows in [('newly_classified_alerts',[r for r in changes if r['gained_any_classification']]),
                       ('lost_classifications',[r for r in changes if r['lost_any_classification']]),
                       ('rule_gain_breakdown',supports),('rule_ablation_comparison',ablations),
                       ('regression_cases',regression),('monthly_coverage',months),('yearly_coverage',years),
                       ('withdrawn_v14_evidence',withdrawn),('pre_alert_only_lexical',pre_only),('all_alert_changes',changes)]:
        write_audit_csv(rows,out/(name+'.csv'))
    with (out/'holdout_manifest.json').open('w',encoding='utf-8') as stream:json.dump(holdout,stream,indent=2)
    holdoutrows=[dict(alert_id=f['alert_id'],v14_categories=f['baseline_categories'],v15_categories=f['threat_combination'],
                      status='UNTOUCHED_DIAGNOSTIC_HOLDOUT_UNLABELED') for f in features if f['alert_id'] in holdout['reserved_alert_ids']]
    write_audit_csv(holdoutrows,out/'holdout_predictions.csv')
    audit=export_audit(root,features,markers,selected,data);audit['strata']=audit_strata
    strength_ids=defaultdict(set)
    for m in markers:
        if m['accepted_for_aggregation']:strength_ids[m['attribution_strength']].add(m['alert_id'])
    gained_uncertain={m['alert_id'] for m in markers if m['accepted_for_aggregation'] and m['uncertainty_flag'] and m['alert_id'] in gained}
    combinations=Counter('+'.join(sorted({cause for m in by[aid] if m['accepted_for_aggregation'] for cause in m['correction_causes']})) for aid in gained)
    report=dict(task_status='COMPLETE_WITH_CONFIRMED_LABEL_DISCREPANCIES' if any(r['regression_status']=='CONFIRMED_LABEL_DISCREPANCY' for r in regression) else 'COMPLETE',
                status=STATUS,v1_3=summary(data['v13']),v1_4=summary(data['v14']),v1_5=summary(features),
                comparison_v14_v15=comparison,comparison_v13_v15=comp13,net_new=comparison['gained_count']-comparison['lost_count'],
                rule_gain_breakdown=supports,rule_ablations=ablations,correction_overlap_combinations=dict(combinations),
                monthly=months,yearly=years,recent_months=[r for r in months if r['month'] in ['2026-08','2026-09','2026-10']],
                regression_cases=regression,holdout=holdout,audit=audit,evidence_strength_alerts={k:len(v) for k,v in strength_ids.items()},
                pre_alert_only_lexical=dict(unique_alerts=len({r['alert_id'] for r in pre_only}),alert_category_pairs=len(pre_only)),
                new_uncertain_classifications=len(gained_uncertain),withdrawn_evidence_records=len(withdrawn),
                protected_hashes=protected,rule_hashes=frozen_rules,input_hashes={k:digest(root/p) for k,p in FILES.items()},
                no_hardcoded_alert_exceptions=True,independent_semantic_validation=False,source_hashes_unchanged=True)
    if snapshot(root)!=protected or rule_hashes(root)!=frozen_rules:
        raise ValueError('Frozen source, classifier, prior report or v1.5 rules changed during run')
    report['elapsed_seconds']=round(time.perf_counter()-tick,6)
    (out/'classification_comparison.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
    (out/'classification_comparison.md').write_text(markdown(report),encoding='utf-8')
    print(json.dumps({k:report[k] for k in ['task_status','net_new','recent_months','rule_gain_breakdown','rule_ablations','pre_alert_only_lexical','new_uncertain_classifications','elapsed_seconds']},ensure_ascii=False))
    print('Counts:',report['v1_5']['classified_alerts'],report['v1_5']['no_marker_with_text'],comparison['gained_count'],comparison['lost_count'])
    return report


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path,default=Path('.'))
    run(parser.parse_args().root)
